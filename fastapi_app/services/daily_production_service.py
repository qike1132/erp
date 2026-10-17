"""
每日生产报表业务逻辑服务

统一三部门（充棉部、手工部、包装部）的 CRUD、每日汇总、月度汇总、导出。
"""
import calendar
import io
import logging
from collections.abc import Iterable
from datetime import date, timedelta
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import and_, func, or_, select, true
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.sql.expression import Join

from fastapi_app.core.business_rules import MAX_QUANTITY
from fastapi_app.core.cache_invalidation import cache_mgr, schedule_post_commit_invalidation
from fastapi_app.core.exception_handler import safe_publish_event
from fastapi_app.core.float_utils import round_money
from fastapi_app.core.production_scope import PRODUCTION_DEPARTMENTS
from fastapi_app.core.sanitizer import escape_like
from fastapi_app.core.validators import validate_not_empty
from fastapi_app.services.notification_service import StockEvent
from fastapi_app.shared.stock_concurrency import (
    apply_stock_delta_locked,
    assert_stock_available,
    get_product_for_update,
)
from fastapi_app.shared_models import Customer, DailyProductionRecord, InventoryLog, Product
from fastapi_app.utils import calc_total_pages, parse_date_with_fallback

logger = logging.getLogger(__name__)

DEPARTMENT_LABELS = {
    'filling': '充棉部',
    'handcraft': '手工部',
    'packaging': '包装部',
}


class InvalidDateRangeError(ValueError):
    """日期区间非法（如结束日期早于开始日期）。

    继承 ValueError，使既有「ValueError → 400」的映射继续生效；
    /export 端点需要把它与「行数超限的 ValueError」区分开（后者映 413）。
    """


async def _apply_production_stock(
    db: AsyncSession, record: DailyProductionRecord, *, delta: int, product_id: int | None = None,
) -> None:
    """生产数量变更同步库存（2026-10-06 用户裁决：生产即入库）。

    对 product.current_stock 施加 delta 并写 InventoryLog 流水；
    负向变化（改小/换产品回退/删除）走全仓统一的负库存准入 assert_stock_available
    ——已出库消耗掉的数量不允许通过改/删生产记录洗成负库存。

    product_id 为空（生产记录未匹配到产品档案）时不产生任何库存变化，
    与 /stock 页 available_stock 口径（production_agg 仅统计 product_id 非空）一致。
    """
    pid = product_id if product_id is not None else record.product_id
    if pid is None or delta == 0:
        return
    product = await get_product_for_update(db, pid)
    if product is None:
        logger.warning('生产记录 %s 库存联动跳过：产品 %s 不存在', record.id, pid)
        return
    if delta < 0:
        assert_stock_available(
            product.current_stock, delta,
            message=f'生产记录变动将导致库存为负数（当前库存：{product.current_stock}，变化：{delta:+d}）',
        )
    # 2026-10-17 审计 F-BE-SVC-STOCK-01：原为 `product.current_stock += delta` 直写，
    # 不递增 version，与走乐观锁的写入口并发时丢失更新。改走统一写入口
    # （调用方已持行锁，apply_stock_delta_locked 内置 WHERE version 守卫与原子非负守卫）。
    await apply_stock_delta_locked(db, product, delta)
    db.add(InventoryLog(
        product_id=pid,
        change_type='生产入库' if delta > 0 else '生产入库冲销',
        change_quantity=delta,
        related_id=str(record.id),
        note=f'生产录入 {record.department} {record.production_date}',
    ))
    # 库存视图（/stock 的「库存」列 = 包装部入库 + 已出货净额）读的是
    # inventory:paginated（ttl=120s / local 60s）：不失效的话，前端收到下面的
    # STOCK_CHANGED 重取也会命中旧缓存，最长 120s 才看到新数量（2026-10-08 实测
    # 「包装录入 5000 → /stock 延迟显示」）。与入库/出库/退货三条写路径同款：
    # 提交前广覆盖 + post-commit 兜底（并发读回写脏值的竞态）。
    await cache_mgr.invalidate_stock_and_order()
    schedule_post_commit_invalidation(db, cache_mgr.invalidate_stock_and_order)
    # 生产即入库也会改变库存视图（库存查询页的生产入库/结余列）——发库存变更事件，
    # 让已接 WebSocket 的页面实时刷新（与入库/出库/退货三条写路径同款）。
    await safe_publish_event(
        stock_service.publish_stock_event, StockEvent.STOCK_CHANGED,
        {
            'product_id': pid,
            'product_name': product.name,
            'change_type': '生产入库' if delta > 0 else '生产入库冲销',
            'quantity': delta,
            'current_stock': product.current_stock,
        },
        context="生产事件发布",
    )


async def _build_product_price_fallback(
    db: AsyncSession,
    records: Iterable[DailyProductionRecord],
) -> dict[int, Decimal]:
    """当记录的 unit_price 为 0 时，按 product_id 回退到产品单价

    返回 {product_id: price} 映射；未匹配到产品的记录保持原值 0

    注：注解为 Decimal 而非 float —— Product.price 是 Numeric 列，
    实现原样返回 Decimal（见下方注释），此前注解写成 float 与实现不符。
    """
    fallback_ids = {r.product_id for r in records if (r.unit_price or 0) == 0 and r.product_id}
    if not fallback_ids:
        return {}
    result = await db.execute(select(Product.id, Product.price).where(Product.id.in_(fallback_ids)))
    # Decimal 原样返回（Product.price 为 Numeric 列），避免 float 中转损耗精度
    return {pid: (price if price is not None else Decimal('0')) for pid, price in result.all()}


# ==================== 校验 ====================

def validate_record_data(
    department: str,
    product_name: str,
    quantity: int,
    unit_price: float = 0,
) -> None:
    """校验生产记录数据"""
    if department not in DEPARTMENT_LABELS:
        raise ValueError(f'部门必须是 filling、handcraft、packaging 之一，当前: {department}')
    validate_not_empty(product_name, "产品名称")
    if quantity <= 0:
        raise ValueError('生产数量必须大于 0')
    if quantity > MAX_QUANTITY:
        raise ValueError(f'生产数量不能超过 {MAX_QUANTITY}')
    if unit_price < 0:
        raise ValueError('单价不能为负数')


def _validate_custom(is_custom: bool | None, custom_country: str | None) -> None:
    """定制标记校验：标了定制就必须有国家（唯一新增校验，见成品库存方案）。"""
    if is_custom and not (custom_country or '').strip():
        raise ValueError('定制产品必须填写定制国家')


async def _resolve_customer_name(db: AsyncSession, customer_id: int | None) -> str | None:
    """校验客户并返回名称快照；空客户直接返回 None。"""
    if customer_id is None:
        return None
    result = await db.execute(
        select(Customer.name).where(Customer.id == customer_id, Customer.status.is_(True))
    )
    customer_name = result.scalar_one_or_none()
    if customer_name is None:
        raise ValueError('客户不存在或已停用')
    return customer_name


# ==================== CRUD ====================

def _build_record_conditions(
    *,
    year: int | None,
    month: int | None,
    start_date: str,
    end_date: str,
    departments: set[str] | None,
    created_by: str,
    production_date: str = '',
    keyword: str = '',
) -> list:
    """构造生产记录查询条件（纯函数，便于单测覆盖组合边界）。

    Args:
        year: 年份；与 month 一起使用。
        month: 月份。
        start_date: 区间起始日 (YYYY-MM-DD)；与 end_date 必须成对，且优先于 year/month。
        end_date: 区间结束日 (YYYY-MM-DD)，闭区间。
        departments: 部门集合；None 表示不限。
        created_by: 录单人；空白视为不限。
        production_date: 指定单日，叠加在日期条件之上。
        keyword: 产品名或 69 码关键词；空白视为不限，LIKE 通配符按字面量转义。

    Returns:
        SQLAlchemy 条件列表。

    Raises:
        ValueError: 未提供任何日期条件、区间只给一半、或区间反向。
    """
    if start_date or end_date:
        if not (start_date and end_date):
            raise ValueError('start_date 与 end_date 必须同时提供')
        range_start = parse_date_with_fallback(start_date)
        range_end = parse_date_with_fallback(end_date)
        if range_end < range_start:
            raise ValueError('结束日期不能早于开始日期')
        conditions: list = [
            DailyProductionRecord.production_date >= range_start,
            DailyProductionRecord.production_date <= range_end,
        ]
    else:
        if year is None or month is None:
            raise ValueError('必须提供 year/month，或同时提供 start_date/end_date')
        month_start = date(year, month, 1)
        month_end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
        conditions = [
            DailyProductionRecord.production_date >= month_start,
            DailyProductionRecord.production_date < month_end,
        ]

    if departments is not None:
        conditions.append(DailyProductionRecord.department.in_(sorted(departments)))

    if created_by and created_by.strip():
        conditions.append(DailyProductionRecord.created_by == created_by.strip())

    if production_date:
        conditions.append(
            DailyProductionRecord.production_date == parse_date_with_fallback(production_date)
        )

    if keyword and keyword.strip():
        pattern = f'%{escape_like(keyword.strip())}%'
        conditions.append(
            or_(
                DailyProductionRecord.product_name.like(pattern, escape='\\'),
                DailyProductionRecord.barcode == keyword.strip(),
            )
        )

    return conditions


def _effective_unit_price_expr():
    """记录单价优先；为 0 时回退到产品当前单价（与行内展示同口径）。

    与 get_monthly_summary 同一惯用法：明细行的 production_value 在
    unit_price=0 时按 Product.price 回退（见 to_dict / _build_product_price_fallback），
    统计 CTE 若直接 sum(记录单价×数量)，页脚合计会**低于**列内各行之和。
    """
    return func.coalesce(
        func.nullif(DailyProductionRecord.unit_price, 0),
        Product.price,
        0,
    )


def _build_stats_cte(conditions: list):
    """构造「单行聚合」统计 CTE，避免与明细行交叉连接放大统计值。

    Args:
        conditions: 与明细查询完全相同的过滤条件。

    Returns:
        stats 子查询的 CTE（total_count/total_qty/total_val/total_sku）。
    """
    return (
        select(
            func.count(DailyProductionRecord.id).label('total_count'),
            func.coalesce(func.sum(DailyProductionRecord.quantity), 0).label('total_qty'),
            func.coalesce(func.sum(_effective_unit_price_expr() * DailyProductionRecord.quantity), 0).label('total_val'),
            func.count(func.distinct(DailyProductionRecord.product_name)).label('total_sku'),
        )
        .outerjoin(Product, DailyProductionRecord.product_id == Product.id)
        .where(*conditions)
        .cte('stats')
    )


def _record_order_by(sort: str) -> tuple:
    """明细行排序键；desc 时同一天内按 id 倒序，避免同日期内顺序随机。"""
    if sort == 'asc':
        return (
            DailyProductionRecord.production_date,
            DailyProductionRecord.department,
            DailyProductionRecord.id,
        )
    return (
        DailyProductionRecord.production_date.desc(),
        DailyProductionRecord.id.desc(),
    )


async def _fetch_records_page(
    db: AsyncSession,
    conditions: list,
    stats_cte,
    sort: str,
    offset: int,
    per_page: int,
) -> list[DailyProductionRecord]:
    """取一页明细行。明细与单行统计 CTE 交叉连接（行 × 统计单行 = 行本身）。"""
    cross_from = Join(DailyProductionRecord.__table__, stats_cte, true())
    result = await db.execute(
        select(
            DailyProductionRecord,
            stats_cte.c.total_count,
            stats_cte.c.total_qty,
            stats_cte.c.total_val,
            stats_cte.c.total_sku,
        )
        .select_from(cross_from)
        .where(*conditions)
        .order_by(*_record_order_by(sort))
        .offset(offset)
        .limit(per_page)
    )
    return [r[0] for r in result.all()]


async def _fetch_records_stats(db: AsyncSession, stats_cte) -> tuple:
    """stats 从 stats_cte 单独取一次，保证与页码无关（越界页 rows 为空时不误报 0）。"""
    stats_row = (await db.execute(
        select(stats_cte.c.total_qty, stats_cte.c.total_val, stats_cte.c.total_sku)
    )).first()
    if not stats_row:
        return 0, 0, 0
    return stats_row[0], stats_row[1], stats_row[2]


async def get_records(
    db: AsyncSession,
    *,
    year: int | None = None,
    month: int | None = None,
    departments: set[str] | None = None,
    start_date: str = '',
    end_date: str = '',
    created_by: str = '',
    sort: str = 'desc',
    page: int = 1,
    per_page: int = 50,
    production_date: str = '',
    keyword: str = '',
) -> dict:
    """获取生产记录列表（分页 + 统计）。

    Args:
        db: 只读会话。
        year / month: 按自然月筛选；提供区间时被忽略。
        departments: 可见部门集合；None 表示全部部门。
        start_date / end_date: 区间筛选（闭区间，可跨月），必须成对。
        created_by: 录单人筛选，空白视为不限。
        sort: 'desc'（默认，最新在前）或 'asc'。
        page: 页码，从 1 开始。
        per_page: 每页数量，服务端收敛到 1..200。
        production_date: 指定单日，叠加在日期条件之上。
        keyword: 产品名或 69 码关键词；空白视为不限，LIKE 通配符按字面量转义。

    Returns:
        {'data', 'total', 'page', 'per_page', 'pages', 'stats'}

    Raises:
        ValueError: 日期条件非法或 sort 取值非法。
    """
    if sort not in ('asc', 'desc'):
        raise ValueError("sort 只能是 'asc' 或 'desc'")

    conditions = _build_record_conditions(
        year=year, month=month, start_date=start_date, end_date=end_date,
        departments=departments, created_by=created_by, production_date=production_date, keyword=keyword,
    )

    per_page = max(1, min(per_page, 200))
    stats_cte = _build_stats_cte(conditions)
    records = await _fetch_records_page(db, conditions, stats_cte, sort, (page - 1) * per_page, per_page)
    total = await _count_records(db, conditions)
    total_qty, total_val, total_sku = await _fetch_records_stats(db, stats_cte)

    price_fallback = await _build_product_price_fallback(db, records)

    return {
        'data': [r.to_dict(price_fallback) for r in records],
        'total': total,
        'page': page,
        'per_page': per_page,
        'pages': calc_total_pages(total, per_page),
        'stats': {
            'total_quantity': int(total_qty),
            'total_value': round_money(float(total_val)),
            'total_sku': int(total_sku),
        },
    }


async def _count_records(db: AsyncSession, conditions: list) -> int:
    """total 用独立 COUNT 查询获取：请求页越界时 rows 为空，此前从 rows[0] 取值会误报 0。"""
    result = await db.execute(select(func.count(DailyProductionRecord.id)).where(*conditions))
    return int(result.scalar() or 0)


async def get_record(db: AsyncSession, record_id: int) -> DailyProductionRecord | None:
    """按 id 取单条生产记录（不存在返回 None，不抛异常）。

    供端点做写权限校验与删除前快照使用。
    """
    result = await db.execute(
        select(DailyProductionRecord).where(DailyProductionRecord.id == record_id)
    )
    return result.scalar_one_or_none()


async def create_record(
    db: AsyncSession,
    *,
    department: str,
    product_id: int | None,
    product_name: str,
    customer_id: int | None = None,
    barcode: str,
    unit_price: float,
    quantity: int,
    production_date_str: str,
    remark: str = '',
    created_by: str = '',
    batch_no: str | None = None,
    is_custom: bool | None = None,
    custom_country: str | None = None,
) -> DailyProductionRecord:
    """创建生产记录"""
    validate_record_data(department, product_name, quantity, unit_price)
    _validate_custom(is_custom, custom_country)
    prod_date = parse_date_with_fallback(production_date_str)

    # 限制不能录入未来日期的数据
    if prod_date > date.today():
        raise ValueError('不能录入未来日期的生产数据')

    customer_name = await _resolve_customer_name(db, customer_id)
    record = DailyProductionRecord(
        department=department,
        product_id=product_id,
        customer_id=customer_id,
        customer_name=customer_name,
        product_name=product_name.strip(),
        barcode=barcode or '',
        unit_price=unit_price,
        quantity=quantity,
        production_date=prod_date,
        remark=remark.strip() if remark else None,
        batch_no=(batch_no or '').strip() or None,
        is_custom=bool(is_custom),
        custom_country=((custom_country or '').strip() or None) if is_custom else None,
        created_by=(created_by or '').strip(),
    )
    db.add(record)
    await db.flush()
    await db.refresh(record)
    # 生产即入库（2026-10-06 用户裁决）：生产数量同步增加可出货库存
    await _apply_production_stock(db, record, delta=quantity)
    logger.info(f'生产记录创建: {department} - {product_name} - {quantity} by {created_by}')
    return record


async def update_record(
    db: AsyncSession,
    record_id: int,
    *,
    created_by: str = '',
    **kwargs,
) -> DailyProductionRecord:
    """更新生产记录"""
    result = await db.execute(
        # F-BE-SVC-STOCK-03：写路径取行锁，delta 不取自陈旧快照
        select(DailyProductionRecord).where(DailyProductionRecord.id == record_id).with_for_update()
    )
    record = result.scalar_one_or_none()
    if not record:
        raise ValueError('记录不存在')

    # 校验更新数据
    dept = kwargs.get('department', record.department)
    name = kwargs.get('product_name', record.product_name)
    qty = kwargs.get('quantity', record.quantity)
    price = kwargs.get('unit_price', record.unit_price)
    validate_record_data(dept, name, qty, price)

    # 库存联动基线（2026-10-06 生产即入库）：变更前快照
    old_product_id, old_quantity = record.product_id, record.quantity
    if 'customer_id' in kwargs:
        record.customer_id = kwargs['customer_id']
        record.customer_name = await _resolve_customer_name(db, kwargs['customer_id'])

    for field in ('department', 'product_id', 'product_name', 'barcode', 'unit_price', 'quantity', 'remark'):
        if field in kwargs and kwargs[field] is not None:
            if field == 'product_name' and isinstance(kwargs[field], str):
                setattr(record, field, kwargs[field].strip())
            elif field == 'remark' and isinstance(kwargs[field], str):
                setattr(record, field, kwargs[field].strip() or None)
            else:
                setattr(record, field, kwargs[field])
    if 'production_date' in kwargs and kwargs['production_date'] is not None:
        new_date = parse_date_with_fallback(kwargs['production_date'])
        if new_date > date.today():
            raise ValueError('不能将生产日期修改为未来日期')
        record.production_date = new_date

    # 批次/定制字段：允许显式置空（判 'in kwargs' 而非 'is not None'）
    if 'batch_no' in kwargs:
        record.batch_no = (kwargs['batch_no'] or '').strip() or None
    if 'is_custom' in kwargs:
        record.is_custom = bool(kwargs['is_custom'])
    if 'custom_country' in kwargs:
        record.custom_country = (kwargs['custom_country'] or '').strip() or None
    # 用最终值校验（改任一字段后 is_custom/custom_country 都可能变化）
    _validate_custom(record.is_custom, record.custom_country)
    if not record.is_custom:
        record.custom_country = None

    await db.flush()

    # 库存联动：同产品按差额调整；换产品则旧产品回退、新产品全额入账
    if old_product_id == record.product_id:
        await _apply_production_stock(db, record, delta=record.quantity - old_quantity)
    else:
        await _apply_production_stock(db, record, delta=-old_quantity, product_id=old_product_id)
        await _apply_production_stock(db, record, delta=record.quantity)

    await db.refresh(record)
    logger.info(f'生产记录更新: {record_id} by {created_by}')
    return record


async def delete_record(db: AsyncSession, record_id: int, *, operator: str = '') -> None:
    """删除生产记录（物理删除）。

    注意：本表没有软删除列，删除不可恢复；审计日志是事后唯一的追溯手段，
    因此**必须在 delete 之前**把完整快照写进日志。

    Args:
        db: 读写会话。
        record_id: 记录 ID。
        operator: 本次操作人用户名。

    Raises:
        ValueError: 记录不存在。
    """
    result = await db.execute(
        # F-BE-SVC-STOCK-03：写路径取行锁，delta 不取自陈旧快照
        select(DailyProductionRecord).where(DailyProductionRecord.id == record_id).with_for_update()
    )
    record = result.scalar_one_or_none()
    if not record:
        raise ValueError('记录不存在')

    logger.info(
        f'生产记录删除: id={record.id} department={record.department} '
        f'product_name={record.product_name} barcode={record.barcode} '
        f'quantity={record.quantity} unit_price={record.unit_price} '
        f'production_date={record.production_date} record_created_by={record.created_by} '
        f'operator={operator}'
    )
    # 库存联动（2026-10-06 生产即入库）：删除生产记录同步回退库存；
    # 若已发生出货导致库存不足，按全仓负库存政策拦截（assert_stock_available）
    await _apply_production_stock(db, record, delta=-record.quantity)
    await db.delete(record)


# ==================== 每日汇总 ====================

async def get_daily_summary(
    db: AsyncSession,
    target_date: str,
    departments: set[str] | None = None,
) -> dict:
    """获取指定日期的各部门汇总

    Args:
        db: 只读会话。
        target_date: 目标日期 (YYYY-MM-DD)。
        departments: 可见部门集合；None 表示全部部门。范围外部门不出现在结果里。
    """
    d = parse_date_with_fallback(target_date)
    dept_keys = (
        list(PRODUCTION_DEPARTMENTS) if departments is None
        else [k for k in PRODUCTION_DEPARTMENTS if k in departments]
    )
    result = await db.execute(
        select(DailyProductionRecord)
        .where(
            DailyProductionRecord.production_date == d,
            DailyProductionRecord.department.in_(dept_keys),
        )
        .order_by(DailyProductionRecord.department, DailyProductionRecord.id)
    )
    records = list(result.scalars().all())
    price_fallback = await _build_product_price_fallback(db, records)

    def effective_unit_price(r: DailyProductionRecord) -> Decimal:
        up = r.unit_price if r.unit_price is not None else Decimal('0')
        if up > 0:
            return up
        if r.product_id and r.product_id in price_fallback:
            return price_fallback[r.product_id]
        return Decimal('0')

    dept_data: dict[str, dict[str, Any]] = {}
    for r in records:
        dept = r.department
        if dept not in dept_data:
            dept_data[dept] = {'quantity': 0, 'value': Decimal('0'), 'skus': set()}
        dept_data[dept]['quantity'] += r.quantity
        dept_data[dept]['value'] += effective_unit_price(r) * r.quantity
        if r.product_name:
            dept_data[dept]['skus'].add(r.product_name)

    department_rows: list[dict[str, Any]] = []
    overall_quantity = 0
    overall_value = Decimal('0')
    overall_skus: set[str] = set()

    for dept_key in dept_keys:
        d_info = dept_data.get(dept_key, {'quantity': 0, 'value': Decimal('0'), 'skus': set()})
        overall_quantity += d_info['quantity']
        overall_value += d_info['value']
        overall_skus.update(d_info['skus'])
        department_rows.append({
            'department': dept_key,
            'department_label': DEPARTMENT_LABELS.get(dept_key, dept_key),
            'total_quantity': d_info['quantity'],
            'total_value': float(round_money(float(d_info['value']))),
            'sku_count': len(d_info['skus']),
        })

    return {
        'date': target_date,
        'departments': department_rows,
        'overall_quantity': overall_quantity,
        'overall_value': float(round_money(float(overall_value))),
        'overall_sku_count': len(overall_skus),
    }


# ==================== 月度汇总 ====================

async def get_monthly_summary(
    db: AsyncSession,
    year: int,
    month: int,
    departments: set[str] | None = None,
) -> dict:
    """获取月度汇总数据

    Args:
        db: 只读会话。
        year: 年份。
        month: 月份。
        departments: 可见部门集合；None 表示全部部门。
            注：daily_data 的每日条目仍保留三个部门键（范围外固定为 0），
            因此每日 total 自动只含范围内部门；monthly_departments 只列范围内部门。
    """
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])
    days_in_month = calendar.monthrange(year, month)[1]
    dept_keys = (
        list(PRODUCTION_DEPARTMENTS) if departments is None
        else [k for k in PRODUCTION_DEPARTMENTS if k in departments]
    )

    # 优先使用记录单价；当记录单价为 0 时回退到产品当前单价
    effective_unit_price = _effective_unit_price_expr()

    # 按日+部门聚合
    result = await db.execute(
        select(
            DailyProductionRecord.production_date,
            DailyProductionRecord.department,
            func.coalesce(func.sum(DailyProductionRecord.quantity), 0).label('total_qty'),
            func.coalesce(func.sum(effective_unit_price * DailyProductionRecord.quantity), 0).label('total_val'),
            func.count(func.distinct(DailyProductionRecord.product_name)).label('sku_count'),
        )
        .outerjoin(Product, DailyProductionRecord.product_id == Product.id)
        .where(and_(
            DailyProductionRecord.production_date >= start,
            DailyProductionRecord.production_date <= end,
            DailyProductionRecord.department.in_(dept_keys),
        ))
        .group_by(DailyProductionRecord.production_date, DailyProductionRecord.department)
        .order_by(DailyProductionRecord.production_date, DailyProductionRecord.department)
    )
    rows = result.all()

    # 组织为 daily_data
    daily_map: dict[str, dict] = {}
    for row in rows:
        d_str = row.production_date.isoformat()
        if d_str not in daily_map:
            daily_map[d_str] = {
                'date': d_str,
                'day': row.production_date.day,
                'filling': {'quantity': 0, 'value': 0.0, 'sku_count': 0},
                'handcraft': {'quantity': 0, 'value': 0.0, 'sku_count': 0},
                'packaging': {'quantity': 0, 'value': 0.0, 'sku_count': 0},
            }
        dept = row.department
        if dept in ('filling', 'handcraft', 'packaging'):
            daily_map[d_str][dept] = {
                'quantity': int(row.total_qty),
                'value': round_money(float(row.total_val)),
                'sku_count': int(row.sku_count),
            }

    daily_data = []
    for day in range(1, days_in_month + 1):
        d_str = date(year, month, day).isoformat()
        entry = daily_map.get(d_str, {
            'date': d_str,
            'day': day,
            'filling': {'quantity': 0, 'value': 0.0, 'sku_count': 0},
            'handcraft': {'quantity': 0, 'value': 0.0, 'sku_count': 0},
            'packaging': {'quantity': 0, 'value': 0.0, 'sku_count': 0},
        })
        entry['day'] = day
        entry['total'] = {
            'quantity': entry['filling']['quantity'] + entry['handcraft']['quantity'] + entry['packaging']['quantity'],
            'value': round_money(entry['filling']['value'] + entry['handcraft']['value'] + entry['packaging']['value']),
            'sku_count': entry['filling']['sku_count'] + entry['handcraft']['sku_count'] + entry['packaging']['sku_count'],
        }
        daily_data.append(entry)

    # 月度部门汇总
    monthly_depts = []
    for dept_key in dept_keys:
        dept_total_qty = sum(d[dept_key]['quantity'] for d in daily_data)
        dept_total_val = round_money(sum(d[dept_key]['value'] for d in daily_data))
        dept_total_sku = sum(d[dept_key]['sku_count'] for d in daily_data)
        monthly_depts.append({
            'department': dept_key,
            'department_label': DEPARTMENT_LABELS.get(dept_key, dept_key),
            'total_quantity': dept_total_qty,
            'total_value': dept_total_val,
            'sku_count': dept_total_sku,
        })

    return {
        'year': year,
        'month': month,
        'daily_data': daily_data,
        'monthly_departments': monthly_depts,
        'monthly_total_quantity': sum(d['total']['quantity'] for d in daily_data),
        'monthly_total_value': round_money(sum(d['total']['value'] for d in daily_data)),
    }


async def get_monthly_summary_for_auto(
    db: AsyncSession,
    year: int,
    month: int,
) -> dict:
    """月度自动汇总（用于定时任务），与 get_monthly_summary 相同逻辑但独立入口便于扩展"""
    return await get_monthly_summary(db, year, month)


# ==================== 趋势 ====================

MAX_TREND_DAYS = 92


async def get_trend(
    db: AsyncSession,
    *,
    start_date: str,
    end_date: str,
    departments: set[str] | None = None,
) -> dict:
    """按日按部门的生产数量趋势（可跨月）。

    Args:
        db: 只读会话。
        start_date: 起始日期 (YYYY-MM-DD)。
        end_date: 结束日期 (YYYY-MM-DD)，闭区间。
        departments: 可见部门集合；None 表示全部部门。

    Returns:
        {'start_date', 'end_date', 'departments', 'days'}；
        days 中每一天都含三个部门键与 total，无数据的日期补 0（让前端 x 轴连续）。

    Raises:
        ValueError: 区间反向或超过 MAX_TREND_DAYS 天。
    """
    start = parse_date_with_fallback(start_date)
    end = parse_date_with_fallback(end_date)
    if end < start:
        raise ValueError('结束日期不能早于开始日期')

    span = (end - start).days + 1
    if span > MAX_TREND_DAYS:
        raise ValueError(f'查询区间不能超过 {MAX_TREND_DAYS} 天')

    dept_keys = (
        list(PRODUCTION_DEPARTMENTS) if departments is None
        else [d for d in PRODUCTION_DEPARTMENTS if d in departments]
    )

    rows: list = []
    if dept_keys:
        result = await db.execute(
            select(
                DailyProductionRecord.production_date,
                DailyProductionRecord.department,
                func.coalesce(func.sum(DailyProductionRecord.quantity), 0).label('qty'),
            )
            .where(and_(
                DailyProductionRecord.production_date >= start,
                DailyProductionRecord.production_date <= end,
                DailyProductionRecord.department.in_(dept_keys),
            ))
            .group_by(DailyProductionRecord.production_date, DailyProductionRecord.department)
        )
        rows = list(result.all())

    by_day: dict[str, dict[str, int]] = {}
    for prod_date, department, qty in rows:
        by_day.setdefault(prod_date.isoformat(), {})[department] = int(qty)

    days = []
    for offset in range(span):
        key = (start + timedelta(days=offset)).isoformat()
        item: dict = {'date': key}
        total = 0
        day_values = by_day.get(key, {})
        for dept in PRODUCTION_DEPARTMENTS:
            value = day_values.get(dept, 0) if dept in dept_keys else 0
            item[dept] = value
            total += value
        item['total'] = total
        days.append(item)

    return {
        'start_date': start_date,
        'end_date': end_date,
        'departments': dept_keys,
        'days': days,
    }


# ==================== 导出 ====================

async def export_records(
    db: AsyncSession,
    *,
    start_date: str = '',
    end_date: str = '',
    departments: set[str] | None = None,
    created_by: str = '',
    keyword: str = '',
) -> bytes:
    """导出生产记录为 Excel（企业标准格式）

    Args:
        db: 只读会话。
        start_date: 区间起始日 (YYYY-MM-DD)。
        end_date: 区间结束日 (YYYY-MM-DD)，闭区间。
        departments: 可见部门集合；None 表示全部部门。
        created_by: 录单人筛选（2026-10-17 审计 F-BE-SVC-STOCK-05：原先导出侧
            自建条件漏掉录单人/关键词，导出的 Excel 与页面所见集合不同）。
        keyword: 产品名或 69 码关键词。
    """
    # 与屏幕侧 get_records 复用同一条件构造器（F-BE-SVC-STOCK-05：
    # 两套实现已漂移——导出丢了 created_by/keyword）。
    try:
        conditions = _build_record_conditions(
            year=None, month=None,
            start_date=start_date, end_date=end_date,
            departments=departments, created_by=created_by, keyword=keyword,
        )
    except InvalidDateRangeError:
        raise
    except ValueError as e:
        # 共享构造器抛普通 ValueError（区间反向等）；导出端点把 ValueError
        # 映射为 413（行数上界），日期错误应保持 400 语义。
        raise InvalidDateRangeError(str(e)) from e

    query = select(DailyProductionRecord).order_by(
        DailyProductionRecord.production_date,
        DailyProductionRecord.department,
        DailyProductionRecord.id,
    )
    if conditions:
        query = query.where(*conditions)

    # 行数上界保护（批次 2）：超限抛 ValueError，端点转 413 引导缩小范围
    from fastapi_app.services.export_limits import ensure_within_export_limit
    await ensure_within_export_limit(db, select(func.count()).select_from(query.subquery()))

    result = await db.execute(query)
    records = list(result.scalars().all())
    price_fallback = await _build_product_price_fallback(db, records)

    def effective_unit_price(r: DailyProductionRecord) -> float:
        up = float(r.unit_price or 0)
        if up > 0:
            return up
        if r.product_id and r.product_id in price_fallback:
            # price_fallback 值为 Decimal（_build_product_price_fallback），统一转 float
            # 与上方 r.unit_price 分支对齐（本函数仅用于 Excel 导出展示，无精度风险）
            return float(price_fallback[r.product_id])
        return 0.0

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '生产报表'

    # 样式定义
    header_fill = PatternFill(start_color='B0E0E6', end_color='B0E0E6', fill_type='solid')
    header_font = Font(bold=True, size=11, color='5A4A4A')
    header_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='D4F1F4'),
        right=Side(style='thin', color='D4F1F4'),
        top=Side(style='thin', color='D4F1F4'),
        bottom=Side(style='thin', color='D4F1F4'),
    )
    summary_fill = PatternFill(start_color='F0F9FF', end_color='F0F9FF', fill_type='solid')
    summary_font = Font(bold=True, size=11, color='1d4ed8')

    # 标题行
    headers = ['生产日期', '部门', '产品名称', '69码', '单价', '数量', '产值', '备注', '创建人', '创建时间']
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # 数据行
    total_qty = 0
    total_val = 0.0
    for r in records:
        up = effective_unit_price(r)
        val = round_money(up * r.quantity)
        total_qty += r.quantity
        total_val += val
        ws.append([
            r.production_date.isoformat() if r.production_date else '',
            DEPARTMENT_LABELS.get(r.department, r.department),
            r.product_name,
            r.barcode or '',
            up,
            r.quantity,
            val,
            r.remark or '',
            r.created_by or '',
            r.create_time.strftime('%Y-%m-%d %H:%M:%S') if r.create_time else '',
        ])

    # 数据行样式
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    # 空数据提示
    if not records:
        ws.append(['暂无数据', '', '', '', '', '', '', '', '', ''])

    # 汇总行
    summary_label = '全部门合计' if departments is None else '合计'
    ws.append([summary_label, '', '', '', '', total_qty, round_money(total_val), '', '', ''])
    for cell in ws[ws.max_row]:
        cell.fill = summary_fill
        cell.font = summary_font
        cell.border = thin_border

    # 按部门汇总
    ws.append([])
    ws.append(['部门汇总', '生产数量', '产值', '', '', '', '', '', '', ''])
    for cell in ws[ws.max_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border

    dept_summary: dict[str, dict] = {}
    for r in records:
        dept = r.department
        if dept not in dept_summary:
            dept_summary[dept] = {'quantity': 0, 'value': 0.0}
        dept_summary[dept]['quantity'] += r.quantity
        dept_summary[dept]['value'] += effective_unit_price(r) * r.quantity

    for dept_key in ('filling', 'handcraft', 'packaging'):
        if dept_key in dept_summary:
            ws.append([
                DEPARTMENT_LABELS.get(dept_key, dept_key),
                dept_summary[dept_key]['quantity'],
                round_money(dept_summary[dept_key]['value']),
                '', '', '', '', '', '', '',
            ])
            for cell in ws[ws.max_row]:
                cell.border = thin_border

    # 列宽
    col_widths = [12, 10, 32, 15, 10, 10, 12, 24, 10, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # 冻结首行
    ws.freeze_panes = 'A2'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ==================== 产品清单集成 ====================

async def get_product_for_production(
    db: AsyncSession,
    product_id: int,
) -> dict | None:
    """获取产品信息用于生产记录（集成产品清单API）"""
    from fastapi_app.shared_models import Product
    result = await db.execute(
        select(Product).where(Product.id == product_id)
    )
    product = result.scalar_one_or_none()
    if not product:
        return None
    return {
        'id': product.id,
        'name': product.name,
        'barcode': product.barcode or '',
        'price': float(product.price) if hasattr(product, 'price') and product.price else 0,
    }


# 底部引入 stock_service 以提供命名空间解析（延迟到调用期，避免循环依赖）
from fastapi_app.services import stock_service
