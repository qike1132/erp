"""
生产计划模块业务逻辑服务

包含手工部录入、包装部录入、汇总动态报表、数据导出等功能。
"""
import calendar
import io
import logging
import uuid
from datetime import date, datetime, timedelta
from typing import Any

import openpyxl
from sqlalchemy import and_, func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from fastapi_app.core.business_time import fmt_business
from fastapi_app.core.cache import cache_result
from fastapi_app.core.cache_invalidation import (
    InvalidationStrategy,
    cache_mgr,
    schedule_post_commit_invalidation,
)
from fastapi_app.core.sanitizer import escape_like
from fastapi_app.shared_models import Case, CaseMovement, CraftWork, PackagingRecord, Product
from fastapi_app.utils import parse_date_with_fallback

# Excel 公式注入清洗（审查 H12）：worker_name/remark/customer/product_name 都是
# 用户可控文本，直入单元格会被 Excel 当公式/DDE 解析。复用权威清洗器，
# 与 purchase_center/export_utils.py 同口径，不另立标准。
from utils.report_exporter import _sanitize_excel_value

logger = logging.getLogger(__name__)

# 导出最大行数（仿 report_data_service），防止 openpyxl 全量内存构建导致 OOM
_MAX_EXPORT_ROWS = 10000

def _last_day_of_month(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


async def _invalidate_craft_cache() -> None:
    """手工部列表缓存失效（写入后立即清，提交后再兜底清一次）。"""
    await cache_mgr.invalidate(InvalidationStrategy.PATTERN, 'production:craft_works:*')

@cache_result(key_prefix="production:craft_works", ttl=60, local_ttl=30,
              key_builder=lambda *args, **kwargs: f"y{kwargs.get('year', args[1] if len(args) > 1 else '')}:m{kwargs.get('month', args[2] if len(args) > 2 else '')}:d{kwargs.get('work_date', args[3] if len(args) > 3 else '')}")
async def get_craft_works(db: AsyncSession, year: int, month: int, work_date: str='') -> dict:
    """
    获取指定月份的手工部生产记录

    返回每条记录 + 按日汇总（含当天皮套出库总数）
    """
    conditions = [CraftWork.work_date >= date(year, month, 1), CraftWork.work_date <= _last_day_of_month(year, month)]
    if work_date:
        conditions.append(CraftWork.work_date == parse_date_with_fallback(work_date))
    result = await db.execute(select(CraftWork).where(*conditions).order_by(CraftWork.work_date, CraftWork.id))
    records = list(result.scalars().all())
    daily_craft: dict[str, int] = {}
    for r in records:
        d = r.work_date.isoformat() if r.work_date else ''
        daily_craft[d] = daily_craft.get(d, 0) + r.quantity
    start = date(year, month, 1)
    end = _last_day_of_month(year, month)
    case_out_daily = await _get_case_stockout_daily(db, start, end)
    daily_summary = {}
    days_in_month = calendar.monthrange(year, month)[1]
    for day in range(1, days_in_month + 1):
        d = date(year, month, day).isoformat()
        daily_summary[d] = {'date': d, 'day': day, 'craft_total': daily_craft.get(d, 0), 'case_out_total': case_out_daily.get(d, 0)}
    return {'data': [r.to_dict() for r in records], 'daily_summary': daily_summary, 'total': len(records)}

async def create_craft_work(db: AsyncSession, *, worker_name: str, quantity: int, work_date_str: str, remark: str='', created_by: str='') -> CraftWork:
    """创建手工部生产记录"""
    work_date_val = parse_date_with_fallback(work_date_str)
    record = CraftWork(worker_name=worker_name, quantity=quantity, work_date=work_date_val, remark=remark, created_by=created_by)
    db.add(record)
    await db.flush()
    await db.refresh(record)
    await _invalidate_craft_cache()
    schedule_post_commit_invalidation(db, _invalidate_craft_cache)
    logger.info(f'手工部记录创建: {worker_name} - {quantity} - {work_date_str}')
    return record

async def update_craft_work(db: AsyncSession, record_id: int, *, created_by: str='', **kwargs) -> CraftWork:
    """更新手工部生产记录"""
    result = await db.execute(select(CraftWork).where(CraftWork.id == record_id))
    record = result.scalar_one_or_none()
    if not record:
        raise ValueError('记录不存在')
    for field in ('worker_name', 'quantity', 'remark'):
        if field in kwargs and kwargs[field] is not None:
            setattr(record, field, kwargs[field])
    if 'work_date' in kwargs and kwargs['work_date'] is not None:
        record.work_date = parse_date_with_fallback(kwargs['work_date'])
    await db.flush()
    await db.refresh(record)
    await _invalidate_craft_cache()
    schedule_post_commit_invalidation(db, _invalidate_craft_cache)
    return record

async def delete_craft_work(db: AsyncSession, record_id: int) -> None:
    """删除手工部生产记录"""
    result = await db.execute(select(CraftWork).where(CraftWork.id == record_id))
    record = result.scalar_one_or_none()
    if not record:
        raise ValueError('记录不存在')
    await db.delete(record)
    await db.flush()
    await _invalidate_craft_cache()
    schedule_post_commit_invalidation(db, _invalidate_craft_cache)

async def export_craft_works(db: AsyncSession, start_date: str='', end_date: str='') -> bytes:
    """导出手工部生产数据为Excel"""
    conditions = []
    if start_date:
        conditions.append(CraftWork.work_date >= parse_date_with_fallback(start_date))
    if end_date:
        conditions.append(CraftWork.work_date <= parse_date_with_fallback(end_date))
    query = select(CraftWork).order_by(CraftWork.work_date, CraftWork.id)
    if conditions:
        query = query.where(*conditions)
    # 导出上界：无日期条件时为全表导出，超限截断并告警（与包装部导出同模式）
    query = query.limit(_MAX_EXPORT_ROWS + 1)
    records = list((await db.scalars(query)).all())
    if len(records) > _MAX_EXPORT_ROWS:
        logger.warning('手工部数据导出超过最大行数 %d，已截断', _MAX_EXPORT_ROWS)
        records = records[:_MAX_EXPORT_ROWS]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '手工部生产数据'
    ws.append(['生产日期', '工人姓名', '生产数量', '备注', '创建时间'])
    daily_total: dict[str, int] = {}
    for r in records:
        d = r.work_date.isoformat() if r.work_date else ''
        daily_total[d] = daily_total.get(d, 0) + r.quantity
        ws.append([d, _sanitize_excel_value(r.worker_name), r.quantity, _sanitize_excel_value(r.remark or ''), fmt_business(r.create_time, '%Y-%m-%d %H:%M:%S')])
    ws.append([])
    ws.append(['日期', '当天生产总数'])
    for d in sorted(daily_total.keys()):
        ws.append([d, daily_total[d]])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

async def get_packaging_records(db: AsyncSession, year: int, month: int, customer: str='', start_date: str='', end_date: str='') -> dict:
    """获取包装部记录列表"""
    conditions = []
    if year and month:
        conditions.extend([PackagingRecord.pack_date >= date(year, month, 1), PackagingRecord.pack_date <= _last_day_of_month(year, month)])
    if customer:
        conditions.append(PackagingRecord.customer.like(f'%{escape_like(customer)}%'))
    if start_date:
        conditions.append(PackagingRecord.pack_date >= parse_date_with_fallback(start_date))
    if end_date:
        conditions.append(PackagingRecord.pack_date <= parse_date_with_fallback(end_date))
    query = select(PackagingRecord).order_by(PackagingRecord.pack_date.desc(), PackagingRecord.id.desc())
    if conditions:
        query = query.where(*conditions)
    result = await db.execute(query)
    records = list(result.scalars().all())
    return {'data': [r.to_dict() for r in records], 'total': len(records)}

async def create_packaging_record(db: AsyncSession, *, customer: str, product_name: str, barcode: str, quantity: int, pack_date_str: str, remark: str='', created_by: str='', upload_batch: str='') -> PackagingRecord:
    """创建包装部记录"""
    pack_date_val = parse_date_with_fallback(pack_date_str)
    record = PackagingRecord(customer=customer, product_name=product_name, barcode=barcode, quantity=quantity, pack_date=pack_date_val, remark=remark, created_by=created_by, upload_batch=upload_batch)
    db.add(record)
    await db.flush()
    await db.refresh(record)
    logger.info(f'包装部记录创建: {customer} - {product_name} - {quantity}')
    return record

async def batch_create_packaging_records(db: AsyncSession, records_data: list[dict], created_by: str='') -> dict:
    """批量创建包装部记录

    每行使用 savepoint（db.begin_nested）隔离，单行 DB 级错误（如约束冲突）
    仅回滚该行，不影响其他行。遵循 batch_import_service 的 savepoint 模式。
    """
    batch_id = str(uuid.uuid4())[:8]
    created_count = 0
    errors = []
    for i, item in enumerate(records_data):
        try:
            async with db.begin_nested():
                customer = str(item.get('customer', '')).strip()
                product_name = str(item.get('product_name', '')).strip()
                barcode = str(item.get('barcode', '')).strip()
                quantity = int(item.get('quantity', 0))
                pack_date_str = str(item.get('pack_date', '')).strip()
                if not customer:
                    raise ValueError('客户为空')
                if not product_name:
                    raise ValueError('产品名称为空')
                if not barcode:
                    raise ValueError('69码为空')
                if quantity <= 0:
                    raise ValueError('数量必须为正数')
                if not pack_date_str:
                    raise ValueError('包装日期为空')
                pack_date_val = parse_date_with_fallback(pack_date_str)
                record = PackagingRecord(customer=customer, product_name=product_name, barcode=barcode, quantity=quantity, pack_date=pack_date_val, remark=str(item.get('remark', '')).strip(), created_by=created_by, upload_batch=batch_id)
                db.add(record)
                await db.flush()
                created_count += 1
        except (ValueError, KeyError, SQLAlchemyError) as e:
            errors.append(f'第{i + 1}行: {e!s}')
    return {'batch_id': batch_id, 'created_count': created_count, 'error_count': len(errors), 'errors': errors}

async def update_packaging_record(db: AsyncSession, record_id: int, *, created_by: str='', **kwargs) -> PackagingRecord:
    """更新包装部记录"""
    result = await db.execute(select(PackagingRecord).where(PackagingRecord.id == record_id))
    record = result.scalar_one_or_none()
    if not record:
        raise ValueError('记录不存在')
    for field in ('customer', 'product_name', 'barcode', 'quantity', 'remark'):
        if field in kwargs and kwargs[field] is not None:
            setattr(record, field, kwargs[field])
    if 'pack_date' in kwargs and kwargs['pack_date'] is not None:
        record.pack_date = parse_date_with_fallback(kwargs['pack_date'])
    await db.flush()
    await db.refresh(record)
    return record

async def delete_packaging_record(db: AsyncSession, record_id: int) -> None:
    """删除包装部记录"""
    result = await db.execute(select(PackagingRecord).where(PackagingRecord.id == record_id))
    record = result.scalar_one_or_none()
    if not record:
        raise ValueError('记录不存在')
    await db.delete(record)
    await db.flush()

def _parse_packaging_rows(file_content: bytes) -> list[tuple]:
    """同步解析包装部 Excel（线程池执行，配合 xlsx_budget 超时——M5）。"""
    from fastapi_app.services.xlsx_budget import (
        MAX_SHEET_COLS,
        MAX_SHEET_ROWS,
        XlsxBudgetError,
        assert_zip_budget,
        clip_iter_rows,
    )

    assert_zip_budget(file_content)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
    except (ValueError, OSError) as e:
        raise ValueError(f'Excel文件解析失败: {e}') from e
    try:
        ws = wb.active
        rows = list(clip_iter_rows(ws))
    finally:
        wb.close()
    # 裁剪多读 1 行：物化后显式拒绝超限（绝不静默丢行/丢列）
    if len(rows) > MAX_SHEET_ROWS:
        raise XlsxBudgetError(f'工作表行数超过上限 {MAX_SHEET_ROWS}')
    if max((len(r) for r in rows), default=0) > MAX_SHEET_COLS:
        raise XlsxBudgetError(f'工作表列数超过上限 {MAX_SHEET_COLS}')
    return rows


async def upload_packaging_excel(db: AsyncSession, file_content: bytes, created_by: str='') -> dict:
    """上传Excel文件，解析并批量创建包装部记录"""
    from fastapi_app.services.xlsx_budget import run_sync_with_timeout

    # M1：专用有界线程池 + 可跟踪终态（不占默认 executor）
    rows = await run_sync_with_timeout(_parse_packaging_rows, file_content)
    if not rows:
        raise ValueError('Excel文件无数据行')
    records_data = []
    errors = []
    preview = []
    for i, row in enumerate(rows, start=2):
        if not row or all(c is None for c in row):
            continue
        try:
            customer = str(row[0] or '').strip()
            product_name = str(row[1] or '').strip()
            barcode = str(row[2] or '').strip()
            quantity = int(row[3]) if row[3] else 0
            pack_date_str = str(row[4] or '').strip()
            remark = str(row[5] or '').strip() if len(row) > 5 and row[5] else ''
            validation_error = ''
            if not customer:
                validation_error = '客户为空'
            elif not product_name:
                validation_error = '产品名称为空'
            elif not barcode:
                validation_error = '69码为空'
            elif quantity <= 0:
                validation_error = '数量必须为正数'
            elif not pack_date_str:
                validation_error = '包装日期为空'
            else:
                try:
                    parse_date_with_fallback(pack_date_str)
                except ValueError as e:
                    validation_error = str(e)
            preview.append({'customer': customer, 'product_name': product_name, 'barcode': barcode, 'quantity': quantity, 'pack_date': pack_date_str, 'remark': remark, 'valid': not bool(validation_error), 'error': validation_error})
            if not validation_error:
                records_data.append({'customer': customer, 'product_name': product_name, 'barcode': barcode, 'quantity': quantity, 'pack_date': pack_date_str, 'remark': remark})
        except (ValueError, IndexError) as e:
            errors.append(f'第{i}行: 数据格式错误 - {e}')
    result = await batch_create_packaging_records(db, records_data, created_by)
    return {'preview': preview, 'total': len(preview), 'valid_count': len(records_data), 'error_count': len(errors) + len([p for p in preview if not p['valid']]), 'created_count': result['created_count'], 'errors': errors + result['errors']}

async def export_packaging_records(db: AsyncSession, start_date: str='', end_date: str='') -> bytes:
    """导出包装部数据为Excel"""
    conditions = []
    if start_date:
        conditions.append(PackagingRecord.pack_date >= parse_date_with_fallback(start_date))
    if end_date:
        conditions.append(PackagingRecord.pack_date <= parse_date_with_fallback(end_date))
    query = select(PackagingRecord).order_by(PackagingRecord.pack_date, PackagingRecord.id)
    if conditions:
        query = query.where(*conditions)
    # 多取 1 条用于判断是否截断，避免全量加载后内存切片
    query = query.limit(_MAX_EXPORT_ROWS + 1)  # sql-safe: escape_like/绑定参数/常量片段
    result = await db.execute(query)
    records = list(result.scalars().all())
    if len(records) > _MAX_EXPORT_ROWS:
        logger.warning('包装部数据导出超过最大行数 %d，已截断', _MAX_EXPORT_ROWS)
        records = records[:_MAX_EXPORT_ROWS]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '包装部生产数据'
    ws.append(['包装日期', '客户', '产品名称', '69码', '数量', '备注', '创建时间'])
    total_qty = 0
    for r in records:
        total_qty += r.quantity
        ws.append([r.pack_date.isoformat() if r.pack_date else '', _sanitize_excel_value(r.customer), _sanitize_excel_value(r.product_name), _sanitize_excel_value(r.barcode), r.quantity, _sanitize_excel_value(r.remark or ''), fmt_business(r.create_time, '%Y-%m-%d %H:%M:%S')])
    ws.append([])
    ws.append(['合计', '', '', '', total_qty])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

async def get_production_summary(db: AsyncSession, year: int, month: int) -> dict:
    """
    汇总动态报表：按日展示 皮套出库 -> 手工部生产 -> 包装部进仓 的流转链路

    字段：年、月、日、配套出库总数、产品名称、产品数量、手工部生产总数、包装部进仓数量
    """
    start = date(year, month, 1)
    end = _last_day_of_month(year, month)
    days_in_month = calendar.monthrange(year, month)[1]
    case_out_data = await _get_case_stockout_detail(db, start, end)
    craft_daily = await _get_craft_work_daily(db, start, end)
    pack_data = await _get_packaging_detail(db, start, end)
    daily_data = []
    for day in range(1, days_in_month + 1):
        d = date(year, month, day).isoformat()
        day_case_outs = case_out_data.get(d, [])
        case_out_total = sum(item['quantity'] for item in day_case_outs)
        craft_total = craft_daily.get(d, 0)
        day_packs = pack_data.get(d, [])
        pack_total = sum(item['quantity'] for item in day_packs)
        product_details = {}
        for item in day_case_outs:
            key = (item['customer'] or '', item['product_name'] or '')
            if key not in product_details:
                product_details[key] = {'customer': item['customer'] or '', 'product_name': item['product_name'] or '', 'case_out_qty': 0, 'pack_qty': 0}
            product_details[key]['case_out_qty'] += item['quantity']
        for item in day_packs:
            key = (item['customer'] or '', item['product_name'] or '')
            if key not in product_details:
                product_details[key] = {'customer': item['customer'] or '', 'product_name': item['product_name'] or '', 'case_out_qty': 0, 'pack_qty': 0}
            product_details[key]['pack_qty'] += item['quantity']
        daily_data.append({'year': year, 'month': month, 'day': day, 'date': d, 'case_out_total': case_out_total, 'craft_total': craft_total, 'pack_total': pack_total, 'products': list(product_details.values())})
    monthly_case_out = sum(d['case_out_total'] for d in daily_data)
    monthly_craft = sum(d['craft_total'] for d in daily_data)
    monthly_pack = sum(d['pack_total'] for d in daily_data)
    monthly_summary = {'case_out_total': monthly_case_out, 'craft_total': monthly_craft, 'pack_total': monthly_pack, 'diff_case_craft': monthly_case_out - monthly_craft, 'diff_craft_pack': monthly_craft - monthly_pack}
    return {'data': daily_data, 'monthly_summary': monthly_summary, 'total': len(daily_data)}

async def _get_case_stockout_daily(db: AsyncSession, start: date, end: date) -> dict:
    """获取皮套出库按日汇总"""
    # 范围条件（>= 当日零点，< 次日零点）而非 func.date()，使 stockout_date 索引可命中
    # 2026-10-03 第八轮审计：原 end_dt = start_dt + timedelta(days=1) 忽略 end 形参，
    # 区间恒为 [start, start+1天) → 每月 2 号起皮套出库汇总恒为 0。改为按 end 计算。
    start_dt = datetime(start.year, start.month, start.day)
    end_dt = datetime(end.year, end.month, end.day) + timedelta(days=1)
    result = await db.execute(select(func.date(CaseMovement.movement_date).label('d'), func.coalesce(func.sum(CaseMovement.quantity), 0).label('total')).where(and_(CaseMovement.direction == 'out', CaseMovement.movement_date >= start_dt, CaseMovement.movement_date < end_dt)).group_by(func.date(CaseMovement.movement_date)))
    return {str(row.d): int(row.total) for row in result.all()}

async def _get_case_stockout_detail(db: AsyncSession, start: date, end: date) -> dict:
    """获取皮套出库按日+客户+产品明细"""
    # 2026-10-03 第八轮审计：同 _get_case_stockout_daily，原区间只覆盖 start 当天。
    start_dt = datetime(start.year, start.month, start.day)
    end_dt = datetime(end.year, end.month, end.day) + timedelta(days=1)
    result = await db.execute(select(func.date(CaseMovement.movement_date).label('d'), CaseMovement.customer, func.coalesce(Product.name, Case.name).label('product_name'), CaseMovement.quantity).join(Case, CaseMovement.case_id == Case.id).outerjoin(Product, Case.product_id == Product.id).where(and_(CaseMovement.direction == 'out', CaseMovement.movement_date >= start_dt, CaseMovement.movement_date < end_dt)).order_by(CaseMovement.movement_date))
    data: dict[str, list[dict[str, Any]]] = {}
    for row in result.all():
        d = str(row.d)
        if d not in data:
            data[d] = []
        data[d].append({'customer': row.customer or '', 'product_name': row.product_name or '', 'quantity': int(row.quantity)})
    return data

async def _get_craft_work_daily(db: AsyncSession, start: date, end: date) -> dict:
    """获取手工部生产按日汇总"""
    result = await db.execute(select(CraftWork.work_date, func.coalesce(func.sum(CraftWork.quantity), 0).label('total')).where(and_(CraftWork.work_date >= start, CraftWork.work_date <= end)).group_by(CraftWork.work_date))
    return {row.work_date.isoformat(): int(row.total) for row in result.all()}

async def _get_packaging_detail(db: AsyncSession, start: date, end: date) -> dict:
    """获取包装部按日+客户+产品明细"""
    result = await db.execute(select(PackagingRecord.pack_date, PackagingRecord.customer, PackagingRecord.product_name, PackagingRecord.quantity).where(and_(PackagingRecord.pack_date >= start, PackagingRecord.pack_date <= end)).order_by(PackagingRecord.pack_date))
    data: dict[str, list[dict[str, Any]]] = {}
    for row in result.all():
        d = row.pack_date.isoformat()
        if d not in data:
            data[d] = []
        data[d].append({'customer': row.customer or '', 'product_name': row.product_name or '', 'quantity': row.quantity})
    return data
