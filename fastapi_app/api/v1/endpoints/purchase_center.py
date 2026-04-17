"""
采购中心 FastAPI 端点

从 Flask app/routes/purchase_center.py 完整迁移，功能完全一致。
使用同步数据库会话 (SyncSession)，因为依赖 Flask 共享数据库。
"""

import io
import os
import zipfile
from datetime import datetime
from decimal import Decimal
from typing import List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
import openpyxl
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session, selectinload

from fastapi_app.core.database import get_sync_db as get_db
from fastapi_app.core.deps import require_permission
from fastapi_app.core.logger import logger
from fastapi_app.services import purchase_center_service as svc

router = APIRouter()


# ============================================================
# Pydantic 请求模型
# ============================================================

class OrderIdsRequest(BaseModel):
    """订单ID列表请求"""
    ids: List[int] = Field(..., min_length=1, description="订单ID列表")

class OrderIdRequest(BaseModel):
    """单个订单ID请求"""
    order_id: int = Field(..., description="订单ID")

class OrderIdsAltRequest(BaseModel):
    """订单ID列表请求(备用字段名)"""
    order_ids: List[int] = Field(..., min_length=1, description="订单ID列表")

class SupplierCreateRequest(BaseModel):
    """创建供应商请求"""
    name: str = Field(..., min_length=1, max_length=100, description="供应商名称")
    contact_person: str = Field('', max_length=50, description="联系人")
    phone: str = Field('', max_length=50, description="电话")
    address: str = Field('', max_length=200, description="地址")
    remark: str = Field('', max_length=500, description="备注")

class SupplierUpdateRequest(BaseModel):
    """更新供应商请求"""
    name: str = Field(..., min_length=1, max_length=100, description="供应商名称")
    contact_person: str = Field('', max_length=50, description="联系人")
    phone: str = Field('', max_length=50, description="电话")
    address: str = Field('', max_length=200, description="地址")
    remark: str = Field('', max_length=500, description="备注")

class CustomerSyncRequest(BaseModel):
    """客户同步请求"""
    customer_ids: Optional[List[int]] = None
    force: bool = False

class CustomerBatchRequest(BaseModel):
    """客户批量查询请求"""
    ids: List[int] = Field(..., min_length=1, description="客户ID列表")

class BomSaveRequest(BaseModel):
    """保存BOM请求"""
    filling: Optional[dict] = None
    fabrics: List[dict] = []
    accessories: List[dict] = []
    packagings: List[dict] = []

class BomCopyRequest(BaseModel):
    """复制BOM请求"""
    source_product_id: int = Field(..., description="源产品ID")
    overwrite: bool = False

class CalculateMaterialRequest(BaseModel):
    """物料计算请求"""
    orders: List[dict] = Field(..., min_length=1, description="订单物料列表")

class PurchaseOrderCreateRequest(BaseModel):
    """创建采购订单请求"""
    materials: List[dict] = Field(..., min_length=1, description="物料列表")
    customer_id: Optional[int] = None
    customer_name: str = ''
    customer_code: str = ''
    remark: str = ''
    unprocessed_order_id: Optional[int] = None

class PurchaseOrderUpdateRequest(BaseModel):
    """更新采购订单请求"""
    status: Optional[str] = None
    remark: Optional[str] = None
    supplier_name: Optional[str] = None
    customer_name: Optional[str] = None
    modify_reason: str = ''

class BatchCompleteRequest(BaseModel):
    """批量完成订单请求"""
    ids: List[int] = Field(..., min_length=1, description="订单ID列表")

class OneClickGenerateRequest(BaseModel):
    """一键生成采购订单请求"""
    order_ids: List[int] = Field(..., min_length=1, description="订单ID列表")

class ExportCompletedRequest(BaseModel):
    """导出已完成订单请求"""
    order_no: Optional[str] = None
    customer_name: Optional[str] = None
    product_name: Optional[str] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None

class FactorySummarySaveRequest(BaseModel):
    """保存工厂订单编辑请求"""
    edits: List[dict] = Field(..., min_length=1, description="编辑列表")

class SummaryPrintRequest(BaseModel):
    """汇总打印请求"""
    ids: List[int] = Field(..., min_length=1, description="订单ID列表")
    format: str = 'print'

class SaveQueryRequest(BaseModel):
    """保存查询请求"""
    name: str = Field(..., min_length=1, description="查询名称")
    query_type: str = 'purchase_order'
    conditions: dict = Field(..., description="查询条件")


# ============================================================
# 产品/BOM API
# ============================================================

@router.get("/api/products")
def api_get_products(
    page: int = Query(1, ge=1),
    per_page: int = Query(50, ge=1, le=500),
    search: str = Query('', max_length=100),
    has_bom: str = Query('all'),
    category: str = Query('', max_length=100),
    sort_by: str = Query('name'),
    sort_order: str = Query('asc'),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_products(
        db, page=page, per_page=per_page, search=search,
        has_bom=has_bom, category=category, sort_by=sort_by, sort_order=sort_order,
    )
    return {'success': True, **result}


@router.get("/api/products/batch-bom")
def api_get_products_batch_bom(
    ids: str = Query(''),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_products_batch_bom(db, ids)
    return {'success': True, **result}


@router.get("/api/products/search-suggestions")
def api_get_search_suggestions(
    q: str = Query(''),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    suggestions = svc.get_search_suggestions(db, q)
    return {'success': True, 'suggestions': suggestions}


@router.get("/api/products/statistics")
def api_get_products_statistics(
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    stats = svc.get_products_statistics(db)
    return {'success': True, 'statistics': stats}


@router.get("/api/product/{product_id}/bom")
def api_get_product_bom(
    product_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_product_bom(db, product_id)
    if not result:
        raise HTTPException(status_code=404, detail="产品不存在")
    return {'success': True, 'product': result, 'materials': result.pop('materials')}


@router.post("/api/product/{product_id}/bom")
def api_save_product_bom(
    product_id: int,
    data: BomSaveRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        msg = svc.save_product_bom(db, product_id, data.model_dump(), _user.username)
        return {'success': True, 'message': msg}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        db.rollback()
        logger.error(f"保存BOM数据失败: {e}")
        raise HTTPException(status_code=500, detail="保存失败，请稍后重试")


@router.post("/api/calculate")
def api_calculate_material(
    data: CalculateMaterialRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.calculate_material(db, data.orders)
        return {'success': True, **result}
    except Exception as e:
        logger.error(f"物料计算失败: {e}")
        raise HTTPException(status_code=500, detail="物料计算失败，请稍后重试")


# ============================================================
# 供应商 API
# ============================================================

@router.get("/api/suppliers")
def api_get_suppliers(
    with_stats: str = Query('0'),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        suppliers = svc.get_suppliers(db, with_stats=(with_stats == '1'))
        return {'success': True, 'suppliers': suppliers}
    except Exception as e:
        logger.exception(f"获取供应商列表失败: {str(e)}")
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


@router.post("/api/suppliers")
def api_create_supplier(
    data: SupplierCreateRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.create_supplier(db, data.model_dump())
        return {'success': True, 'supplier': result}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        db.rollback()
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


@router.put("/api/suppliers/{supplier_id}")
def api_update_supplier(
    supplier_id: int,
    data: SupplierUpdateRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.update_supplier(db, supplier_id, data.model_dump())
        return {'success': True, 'supplier': result}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        db.rollback()
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


@router.get("/api/suppliers/{supplier_id}/materials")
def api_get_supplier_materials(
    supplier_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        materials = svc.get_supplier_materials(db, supplier_id)
        return {'success': True, 'materials': materials}
    except Exception as e:
        logger.exception(f"获取供应商物料失败: {str(e)}")
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


# ============================================================
# 客户管理 API
# ============================================================

@router.get("/api/customers")
def api_get_customers(
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    status: str = Query('all'),
    keyword: str = Query(''),
    sort_by: str = Query('name'),
    sort_order: str = Query('asc'),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_customers(
        db, page=page, per_page=min(per_page, 100), status=status,
        keyword=keyword, sort_by=sort_by, sort_order=sort_order,
    )
    return {'success': True, **result}


@router.post("/api/customers/sync")
def api_sync_customers(
    data: CustomerSyncRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    customer_ids = data.customer_ids
    force = data.force

    from fastapi_app.shared_models import Customer
    customers = db.query(Customer)
    if customer_ids:
        customers = customers.filter(Customer.id.in_(customer_ids))
    customers = customers.all()

    synced_count = 0
    skipped_count = 0
    details = []

    for customer in customers:
        try:
            synced_count += 1
            details.append({
                'id': customer.id, 'name': customer.name,
                'status': 'success', 'sync_time': datetime.now().isoformat(),
            })
        except Exception:
            skipped_count += 1
            details.append({
                'id': customer.id, 'name': customer.name,
                'status': 'failed', 'error': '同步失败',
            })

    return {
        'success': True,
        'message': f'同步完成: 成功{synced_count}条, 跳过{skipped_count}条',
        'synced_count': synced_count,
        'skipped_count': skipped_count,
        'details': details,
    }


@router.post("/api/customers/batch")
def api_batch_get_customers(
    data: CustomerBatchRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    customers = svc.batch_get_customers(db, data.ids)
    return {'success': True, 'customers': customers}


@router.get("/api/customers/statistics")
def api_get_customers_statistics(
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    stats = svc.get_customers_statistics(db)
    return {'success': True, 'statistics': stats}


@router.get("/api/customers/{customer_id}")
def api_get_customer_detail(
    customer_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_customer_detail(db, customer_id)
    if not result:
        raise HTTPException(status_code=404, detail="客户不存在")
    return {'success': True, 'customer': result}


# ============================================================
# 物料分类 & 缓存
# ============================================================

@router.get("/api/categories")
async def api_get_material_categories(
    _user=Depends(require_permission("product_manage")),
):
    from fastapi_app.shared_models import MATERIAL_CATEGORIES
    return {'success': True, 'categories': MATERIAL_CATEGORIES}


@router.post("/api/cache/clear")
async def api_clear_cache(
    _user=Depends(require_permission("product_manage")),
):
    return {'success': True, 'message': '缓存已清除'}


# ============================================================
# 采购订单 API
# ============================================================

@router.post("/api/orders")
def api_create_purchase_order(
    data: PurchaseOrderCreateRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.create_purchase_order(db, data.model_dump(), _user.id, _user.username)
        return {'success': True, **result}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        db.rollback()
        logger.error(f"创建采购订单失败: {e}")
        raise HTTPException(status_code=500, detail="创建订单失败，请稍后重试")


@router.post("/api/orders/one-click-generate")
def api_one_click_generate_order(
    data: OneClickGenerateRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.one_click_generate_order(db, data.order_ids, _user.id, _user.username)
        return {'success': True, **result}
    except Exception as e:
        db.rollback()
        logger.exception(f"一键生成采购订单失败: {str(e)}")
        raise HTTPException(status_code=500, detail="生成采购订单失败，请稍后重试")


@router.get("/api/orders")
def api_get_purchase_orders(
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1),
    status: str = Query('', max_length=50),
    customer: str = Query('', max_length=100),
    product_name: str = Query('', max_length=100),
    start_date: str = Query('', max_length=20),
    end_date: str = Query('', max_length=20),
    order_no: str = Query('', max_length=50),
    supplier: str = Query('', max_length=100),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_purchase_orders(
        db, page=page, per_page=per_page, status=status,
        customer=customer, product_name=product_name,
        start_date=start_date, end_date=end_date,
        order_no=order_no, supplier=supplier,
    )
    return {'success': True, **result}


# 已完成订单相关路由（必须在 /api/orders/{order_id} 之前，避免路径冲突）

@router.get("/api/orders/completed")
def api_get_completed_purchase_orders(
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1),
    order_no: str = Query('', max_length=50),
    customer_name: str = Query('', max_length=100),
    product_name: str = Query('', max_length=100),
    supplier_name: str = Query('', max_length=100),
    barcode: str = Query('', max_length=50),
    material_name: str = Query('', max_length=100),
    start_date: str = Query('', max_length=20),
    end_date: str = Query('', max_length=20),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_completed_purchase_orders(
        db, page=page, per_page=per_page, order_no=order_no,
        customer_name=customer_name, product_name=product_name,
        supplier_name=supplier_name, barcode=barcode,
        material_name=material_name,
        start_date=start_date, end_date=end_date,
    )
    return {'success': True, **result}


@router.post("/api/orders/completed/export")
def api_export_completed_orders(
    data: ExportCompletedRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    html = svc.export_completed_orders(db, data.model_dump(exclude_none=True))
    return {'success': True, 'html': html}


@router.get("/api/orders/completed/suggestions")
def api_get_completed_order_suggestions(
    field: str = Query(''),
    q: str = Query(''),
    limit: int = Query(10, ge=1, le=50),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    suggestions = svc.get_completed_order_suggestions(db, field, q, limit)
    return {'success': True, 'suggestions': suggestions}


@router.get("/api/orders/completed/quick-filters")
async def api_get_completed_order_quick_filters(
    _user=Depends(require_permission("product_manage")),
):
    quick_filters = svc.get_completed_order_quick_filters()
    return {'success': True, 'quick_filters': quick_filters}


@router.get("/api/orders/{order_id}")
def api_get_purchase_order(
    order_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_purchase_order(db, order_id)
    if not result:
        raise HTTPException(status_code=404, detail="订单不存在")
    return {'success': True, 'order': result}


@router.put("/api/orders/{order_id}")
def api_update_purchase_order(
    order_id: int,
    data: PurchaseOrderUpdateRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.update_purchase_order(db, order_id, data.model_dump(exclude_unset=True), _user.username)
    if not result:
        raise HTTPException(status_code=404, detail="订单不存在")
    return {'success': True, 'order': result}


@router.delete("/api/orders/{order_id}")
def api_delete_purchase_order(
    order_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.delete_purchase_order(db, order_id)
    if not result:
        raise HTTPException(status_code=404, detail="订单不存在")
    return {'success': True, 'message': result}


@router.post("/api/orders/complete")
def api_complete_purchase_order(
    data: OrderIdRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.complete_purchase_order(db, data.order_id, _user.id, _user.username)
        return {'success': True, **result}
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        db.rollback()
        logger.error(f"完成订单失败: {e}")
        raise HTTPException(status_code=500, detail="完成订单失败，请稍后重试")


@router.post("/api/orders/batch-complete")
def api_batch_complete_purchase_orders(
    data: BatchCompleteRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.batch_complete_purchase_orders(db, data.ids, _user.id, _user.username)
        return {'success': True, **result}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail="完成订单失败，请稍后重试")


@router.get("/api/orders/{order_id}/print")
def api_get_order_print_data(
    order_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_purchase_order(db, order_id)
    if not result:
        raise HTTPException(status_code=404, detail="订单不存在")
    return {'success': True, 'order': result}


@router.post("/api/orders/batch-print")
def api_get_orders_batch_print_data(
    data: OrderIdsRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    from fastapi_app.shared_models import PurchaseOrder
    orders = db.query(PurchaseOrder).filter(PurchaseOrder.id.in_(data.ids)).all()
    return {'success': True, 'orders': [o.to_dict() for o in orders]}


# ============================================================
# 未处理订单 API
# ============================================================

@router.get("/api/unprocessed-orders")
def api_get_unprocessed_orders(
    customer: str = Query(''),
    page: int = Query(1, ge=1),
    per_page: int = Query(50, ge=1),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_unprocessed_orders(db, customer=customer, page=page, per_page=per_page)
    return {'success': True, **result}


# ============================================================
# 修改历史 API
# ============================================================

@router.get("/api/modification-logs")
def api_get_modification_logs(
    target_type: str = Query(''),
    target_id: int = Query(0),
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_modification_logs(
        db, target_type=target_type, target_id=target_id,
        page=page, per_page=per_page,
    )
    return {'success': True, **result}


# ============================================================
# 保存查询 API
# ============================================================

@router.get("/api/saved-queries")
def api_get_saved_queries(
    query_type: str = Query('purchase_order'),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    queries = svc.get_saved_queries(db, query_type=query_type, user_id=_user.id)
    return {'success': True, 'queries': queries}


@router.post("/api/saved-queries")
def api_save_query(
    data: SaveQueryRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.save_query(db, data.model_dump(), _user.id, _user.username)
        return {'success': True, 'query': result}
    except Exception:
        db.rollback()
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


@router.delete("/api/saved-queries/{query_id}")
def api_delete_saved_query(
    query_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.delete_saved_query(db, query_id)
    if not result:
        raise HTTPException(status_code=404, detail="查询不存在")
    return {'success': True, 'message': result}


@router.post("/api/saved-queries/{query_id}/use")
def api_use_saved_query(
    query_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.use_saved_query(db, query_id)
    if not result:
        raise HTTPException(status_code=404, detail="查询不存在")
    return {'success': True, 'query': result}


# ============================================================
# 订单状态日志 API
# ============================================================

@router.get("/api/status-logs")
def api_get_status_logs(
    target_type: str = Query(''),
    target_id: int = Query(0),
    order_no: str = Query(''),
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    result = svc.get_status_logs(
        db, target_type=target_type, target_id=target_id,
        order_no=order_no, page=page, per_page=per_page,
    )
    return {'success': True, **result}


# ============================================================
# 工厂内部采购订单 API
# ============================================================

@router.post("/api/orders/factory-summary")
def api_get_factory_order_summary(
    data: OrderIdsRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.get_factory_order_summary(db, data.ids)
        return {'success': True, **result}
    except Exception as e:
        logger.exception(f"获取工厂订单汇总失败: {str(e)}")
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


@router.post("/api/orders/generate-factory")
async def api_generate_factory_order(
    data: dict,
    _user=Depends(require_permission("product_manage")),
):
    return {'success': True, 'message': '工厂内部采购订单生成成功（演示版本）'}


@router.post("/api/orders/factory-summary/save")
def api_save_factory_summary_edits(
    data: FactorySummarySaveRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.save_factory_summary_edits(db, data.edits)
        return {'success': True, **result}
    except Exception as e:
        db.rollback()
        logger.exception(f"保存工厂订单编辑失败: {str(e)}")
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


# ============================================================
# 供应商采购订单汇总 API
# ============================================================

@router.post("/api/orders/supplier-summary")
def api_get_supplier_order_summary(
    data: OrderIdsRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    try:
        result = svc.get_supplier_order_summary(db, data.ids)
        return {'success': True, **result}
    except Exception as e:
        logger.exception(f"获取供应商订单汇总失败: {str(e)}")
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


@router.post("/api/orders/supplier-summary-by-vendor")
def api_get_supplier_summary_by_vendor(
    data: OrderIdsRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    """按供应商维度汇总采购订单"""
    try:
        sorted_suppliers = svc.query_supplier_groups(db, data.ids)
        return {
            'success': True,
            'suppliers': sorted_suppliers,
            'supplier_count': len(sorted_suppliers),
            'order_count': len(data.ids),
        }
    except Exception as e:
        logger.exception(f"按供应商汇总失败: {str(e)}")
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


@router.post("/api/orders/generate-supplier")
async def api_generate_supplier_order(
    data: dict,
    _user=Depends(require_permission("product_manage")),
):
    return {'success': True, 'message': '供应商采购订单生成成功（演示版本）'}




# ============================================================
# 汇总打印 API
# ============================================================

@router.post("/api/orders/summary-print")
def api_summary_print(
    data: SummaryPrintRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    ids = data.ids
    format_type = data.format
    try:
        from fastapi_app.shared_models import PurchaseOrder, PurchaseOrderItem, Order

        orders = db.query(PurchaseOrder).filter(PurchaseOrder.id.in_(ids)).options(selectinload(PurchaseOrder.items)).all()

        if format_type in ['factory', 'supplier']:
            source_order_ids = list(set([o.source_order_id for o in orders if o.source_order_id]))
            if source_order_ids:
                related = db.query(PurchaseOrder).filter(PurchaseOrder.source_order_id.in_(source_order_ids)).options(selectinload(PurchaseOrder.items)).all()
                all_order_ids = list(set(ids) | set(o.id for o in related))
            else:
                all_order_ids = list(ids)
            if len(all_order_ids) != len(ids):
                orders = db.query(PurchaseOrder).filter(PurchaseOrder.id.in_(all_order_ids)).options(selectinload(PurchaseOrder.items)).all()
                ids = all_order_ids

        items = db.query(PurchaseOrderItem).filter(PurchaseOrderItem.order_id.in_(ids)).all()
        order_items_map = {}
        for item in items:
            order_items_map.setdefault(item.order_id, []).append(item)

        if format_type == 'factory':
            html = _generate_factory_summary_print(db, orders, order_items_map)
            return {'success': True, 'html': html, 'order_count': len(orders)}
        elif format_type == 'supplier':
            sorted_suppliers = svc.query_supplier_groups(db, ids)
            html = _generate_supplier_summary_print(sorted_suppliers, orders)
            return {'success': True, 'html': html, 'supplier_count': len(sorted_suppliers), 'order_count': len(orders)}
        else:
            html = _generate_normal_summary_print(orders, order_items_map)
            return {'success': True, 'html': html, 'order_count': len(orders)}
    except Exception as e:
        logger.exception(f"汇总打印失败: {str(e)}")
        raise HTTPException(status_code=500, detail="操作失败，请稍后重试")


def _build_print_html(title: str, body_content: str, extra_css: str = "") -> str:
    """生成打印HTML页面的通用模板

    Args:
        title: 页面标题
        body_content: body内的HTML内容
        extra_css: 额外的CSS样式
    """
    base_css = """body{font-family:"Microsoft YaHei",sans-serif;padding:20px}
.page-break{page-break-after:always}
.header{text-align:center;margin-bottom:20px;border-bottom:2px solid #333;padding-bottom:10px}
.header h1{margin:0;color:#333}
.print-date{text-align:right;color:#666;font-size:12px;margin-top:20px}
"""
    return f'''<!DOCTYPE html><html><head><meta charset="utf-8"><title>{title}</title>
<style>{base_css}{extra_css}</style></head><body>
{body_content}
<div class="print-date">打印时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</div></body></html>'''


def _generate_normal_summary_print(orders, order_items_map):
    """生成普通格式打印HTML"""
    extra_css = """.order-section{margin-bottom:30px}
.order-title{font-size:16px;font-weight:bold;margin:15px 0 10px;padding:8px;background:#f5f5f5}
table{width:100%;border-collapse:collapse;margin-top:10px}
th,td{border:1px solid #ddd;padding:8px;text-align:left;font-size:12px}
th{background:#f5f5f5;font-weight:bold}
.order-meta{margin-bottom:10px;font-size:13px}"""
    body_parts = ['<div class="header"><h1>东莞市滨发玩具有限公司</h1><h2>采购订单汇总</h2></div>']
    for idx, order in enumerate(orders):
        if idx > 0:
            body_parts.append('<div class="page-break"></div>')
        items = order_items_map.get(order.id, [])
        body_parts.append(f'<div class="order-section"><div class="order-title">订单编号: {order.order_no}</div>')
        body_parts.append(f'<div class="order-meta"><strong>供应商:</strong> {order.supplier_name or "-"} | <strong>客户:</strong> {order.customer_name or "-"} | <strong>日期:</strong> {order.order_date or "-"}</div>')
        body_parts.append('<table><thead><tr><th>物料名称</th><th>规格</th><th>数量</th><th>单位</th><th>分类</th></tr></thead><tbody>')
        for item in items:
            body_parts.append(f'<tr><td>{item.material_name}</td><td>{item.specification or "-"}</td><td>{item.quantity}</td><td>{item.unit or "-"}</td><td>{item.material_category or "-"}</td></tr>')
        body_parts.append('</tbody></table></div>')
    return _build_print_html("采购订单汇总", ''.join(body_parts), extra_css)


def _generate_factory_summary_print(db, orders, order_items_map):
    """生成工厂内部汇总格式打印HTML"""
    from fastapi_app.shared_models import Order
    source_order_ids = [o.source_order_id for o in orders if o.source_order_id]
    source_orders = db.query(Order).filter(Order.id.in_(source_order_ids)).all() if source_order_ids else []
    source_order_map = {o.id: o for o in source_orders}

    grouped_data = {}
    order_count_map = {}

    for order in orders:
        items = order_items_map.get(order.id, [])
        source_order = source_order_map.get(order.source_order_id) if order.source_order_id else None
        product_name = order.product_name
        if not product_name and source_order and source_order.product:
            product_name = source_order.product.name
        elif not product_name:
            product_name = order.customer_name or '未指定产品'
        product_qty = order.product_quantity
        if not product_qty and source_order:
            product_qty = source_order.quantity
        customer_name = source_order.customer if source_order else (order.customer_name or '未指定客户')
        cp_key = (customer_name, product_name)
        if cp_key not in grouped_data:
            grouped_data[cp_key] = {}
            order_count_map[cp_key] = {'order_count': 0, 'total_fabric_qty': 0, 'product_quantity': product_qty}
        for item in items:
            supplier_name = item.supplier_name or order.supplier_name or '未指定供应商'
            fabric_name = item.material_name or '未指定布料'
            fabric_spec = item.specification or ''
            fabric_qty = item.quantity or 0
            if supplier_name not in grouped_data[cp_key]:
                grouped_data[cp_key][supplier_name] = {}
            fabric_key = (fabric_name, fabric_spec, item.unit or '')
            if fabric_key not in grouped_data[cp_key][supplier_name]:
                grouped_data[cp_key][supplier_name][fabric_key] = {'fabric_name': fabric_name, 'specification': fabric_spec, 'quantity': fabric_qty, 'unit': item.unit or ''}
            else:
                existing = grouped_data[cp_key][supplier_name][fabric_key]
                existing['quantity'] = (float(existing['quantity']) if existing['quantity'] else 0) + (float(fabric_qty) if fabric_qty else 0)
            order_count_map[cp_key]['total_fabric_qty'] += fabric_qty

    for order in orders:
        source_order = source_order_map.get(order.source_order_id) if order.source_order_id else None
        product_name = order.product_name
        if not product_name and source_order and source_order.product:
            product_name = source_order.product.name
        elif not product_name:
            product_name = order.customer_name or '未指定产品'
        customer_name = source_order.customer if source_order else (order.customer_name or '未指定客户')
        cp_key = (customer_name, product_name)
        if cp_key in order_count_map:
            order_count_map[cp_key]['order_count'] += 1

    sorted_groups = {}
    for (customer_name, product_name), supplier_dict in grouped_data.items():
        suppliers_list = []
        for supplier_name, fabric_dict in supplier_dict.items():
            fabrics_list = list(fabric_dict.values())
            for f in fabrics_list:
                f['quantity'] = round(float(f['quantity']), 2) if f['quantity'] else 0
            suppliers_list.append({'supplier_name': supplier_name, 'fabrics': fabrics_list})
        sorted_groups[(customer_name, product_name)] = suppliers_list

    extra_css = """@page{margin:0in}body{margin:0;padding:0}
.header h1{font-size:24pt;font-weight:bold}.header h2{margin:5px 0 0 0;font-size:18pt;font-weight:bold}
table{width:850.15pt;border-collapse:collapse;table-layout:fixed}
th,td{border:1.0pt solid #000000;padding:1px;text-align:center;vertical-align:middle;font-size:10.5pt;white-space:normal}
th{background:#E6E6E6;font-weight:bold}
.total-row{background-color:#F0F0F0;font-weight:bold}"""
    body_parts = ['<div class="header"><h1>东莞市滨发玩具有限公司</h1><h2>工厂内部采购订单汇总</h2></div>']
    body_parts.append('<table><thead><tr><th>客户</th><th>产品名称</th><th>订单数量</th><th>供应商</th><th>布料名称</th><th>规格型号</th><th>数量</th></tr></thead><tbody>')

    grand_total = 0
    for (customer_name, product_name), suppliers in sorted(sorted_groups.items()):
        total_fabric_rows = sum(len(s['fabrics']) for s in suppliers)
        order_info = order_count_map[(customer_name, product_name)]
        global_fabric_idx = 0
        for supplier in suppliers:
            for idx, fabric in enumerate(supplier['fabrics']):
                body_parts.append('<tr>')
                if global_fabric_idx == 0:
                    body_parts.append(f'<td rowspan="{total_fabric_rows}" style="font-weight:bold">{customer_name}</td>')
                    body_parts.append(f'<td rowspan="{total_fabric_rows}">{product_name}</td>')
                    body_parts.append(f'<td rowspan="{total_fabric_rows}">{float(order_info["product_quantity"]) if order_info["product_quantity"] else 0:g}</td>')
                if idx == 0:
                    body_parts.append(f'<td rowspan="{len(supplier["fabrics"])}">{supplier["supplier_name"]}</td>')
                body_parts.append(f'<td style="text-align:left">{fabric["fabric_name"]}</td>')
                body_parts.append(f'<td style="text-align:left">{fabric.get("specification", "")}</td>')
                body_parts.append(f'<td style="text-align:right">{fabric["quantity"]} {fabric["unit"]}</td>')
                body_parts.append('</tr>')
                grand_total += float(fabric["quantity"]) if fabric["quantity"] else 0
                global_fabric_idx += 1

    body_parts.append(f'<tfoot><tr class="total-row"><td colspan="4">合计</td><td></td><td></td><td style="text-align:right">{grand_total}</td></tr></tfoot>')
    body_parts.append('</table>')
    return _build_print_html("工厂内部采购订单汇总", ''.join(body_parts), extra_css)


def _generate_supplier_summary_print(supplier_groups, orders):
    """生成供应商汇总打印HTML"""
    extra_css = """@media print{@page{size:A4;margin:15mm}.no-print{display:none!important}}
.company-header h1{font-size:24px;font-weight:bold}.company-header h2{margin:10px 0 0 0;color:#666;font-size:18px}
.summary-info{margin:15px 0;padding:10px;background:#f9f9f9;border:1px solid #ddd}
.supplier-section{margin-bottom:30px;border:1px solid #ccc}
.supplier-header{background:#4a90e2;color:white;padding:12px 15px;font-size:16px;font-weight:bold}
.supplier-info{padding:10px 15px;background:#f5f5f5;border-bottom:1px solid #ddd}
table{width:100%;border-collapse:collapse}
th,td{border:1px solid #ddd;padding:8px;text-align:left;vertical-align:middle}
th{background:#e8e8e8;font-weight:bold;text-align:center}
.total-row{background:#fff3cd;font-weight:bold}
.footer{margin-top:30px;text-align:right;color:#666;font-size:11px}"""
    body_parts = [f'<div class="company-header"><h1>东莞市滨发玩具有限公司</h1><h2>供应商采购订单汇总</h2></div>']
    body_parts.append(f'<div class="summary-info"><strong>汇总时间：</strong>{datetime.now().strftime("%Y-%m-%d %H:%M:%S")} | <strong>供应商数量：</strong>{len(supplier_groups)} | <strong>订单数量：</strong>{len(orders)}</div>')
    for idx, supplier in enumerate(supplier_groups):
        if idx > 0:
            body_parts.append('<div class="page-break"></div>')
        total_qty = sum(float(item['quantity']) for item in supplier['items'])
        body_parts.append(f'<div class="supplier-section"><div class="supplier-header">供应商：{supplier["supplier_name"]}</div>')
        body_parts.append(f'<div class="supplier-info"><strong>关联订单：</strong>{", ".join(supplier["orders"])} | <strong>物料种类：</strong>{len(supplier["items"])} 种</div>')
        body_parts.append('<table><thead><tr><th>序号</th><th>商品名称</th><th>规格型号</th><th>单位</th><th>数量</th><th>物料分类</th><th>关联订单</th></tr></thead><tbody>')
        for item_idx, item in enumerate(supplier['items'], 1):
            body_parts.append(f'<tr><td style="text-align:center">{item_idx}</td><td>{item["material_name"]}</td><td>{item["specification"] or "-"}</td><td style="text-align:center">{item["unit"] or "-"}</td><td style="text-align:right">{float(item["quantity"]):,.2f}</td><td style="text-align:center">{item["material_category"] or "-"}</td><td>共{item["order_count"]}单</td></tr>')
        body_parts.append(f'<tfoot><tr class="total-row"><td colspan="4" style="text-align:right">合计：</td><td style="text-align:right">{total_qty:,.2f}</td><td colspan="2"></td></tr></tfoot></table></div>')
    return _build_print_html("供应商采购订单汇总", ''.join(body_parts), extra_css)


# ============================================================
# 供应商采购订单 Excel 导出（基于模板填充）
# ============================================================
# 模板文件路径（.xlsx 格式，由 .xls 模板转换而来，存放于项目 static/files/ 目录）
_PROJECT_ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..'))
_TEMPLATE_PATH = os.path.join(_PROJECT_ROOT, 'static', 'files', 'purchase_order_template.xlsx')


def _validate_template_path(template_path: str) -> str:
    """验证模板文件路径在项目目录内，防止路径遍历"""
    normalized = os.path.normpath(template_path)
    if not normalized.startswith(_PROJECT_ROOT):
        raise ValueError("模板路径超出项目目录范围")
    if not os.path.isfile(normalized):
        raise FileNotFoundError(f"模板文件不存在: {normalized}")
    return normalized

# 模板中动态字段对应的单元格位置（openpyxl 1-indexed）
# R4C1  = 供应商名称:xxx
# R10C1~R15C1 = 序号 (数据行1-6)
# R10C3~R15C3 = 品名/规格/色号
# R10C14~R15C14 = 订购数量
# R10C22~R15C22 = 单价/码
# R10C34~R15C34 = 交货日期
# R16C30 = 总金额
# R18C30 = 应付金额


def _copy_font_style(font):
    """复制 Font 对象"""
    from openpyxl.styles import Font as Ft
    return Ft(name=font.name, size=font.size, bold=font.bold, italic=font.italic,
              color=font.color, underline=font.underline, strike=font.strike)


def _copy_align_style(alignment):
    """复制 Alignment 对象"""
    from openpyxl.styles import Alignment as Al
    return Al(horizontal=alignment.horizontal, vertical=alignment.vertical,
              wrap_text=alignment.wrap_text, text_rotation=alignment.text_rotation,
              indent=alignment.indent, shrink_to_fit=alignment.shrink_to_fit)


def _copy_border_style(border):
    """复制 Border 对象"""
    from openpyxl.styles import Border as Bd, Side as Sd
    def _cp(side):
        return Sd(style=side.style, color=side.color) if side.style else Sd()
    return Bd(left=_cp(border.left), right=_cp(border.right),
              top=_cp(border.top), bottom=_cp(border.bottom))


def _fill_cell(ws, row, col, value, number_format=None):
    """向模板中的指定单元格填充值，保留原有格式"""
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if number_format:
        cell.number_format = number_format


def create_supplier_workbook(supplier: dict) -> Workbook:
    """基于模板文件为单个供应商创建采购订单 Excel 工作簿

    策略：加载预设模板 → 仅替换动态字段 → 保留模板所有格式不变
    支持动态行数：物料数超过模板固定6行时自动插入新行，不足6行补空白行
    """
    template_path = os.path.normpath(_TEMPLATE_PATH)
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    supplier_name = supplier['supplier_name']
    items = supplier['items']
    supplier_address = supplier.get('address', '') or ''
    supplier_phone = supplier.get('phone', '') or ''
    supplier_contact = supplier.get('contact_person', '') or ''

    # ── 1. 供应商名称（R4C1: "供应商名称:xxx"） ──
    _fill_cell(ws, 4, 1, f'供应商名称:{supplier_name}')

    # ── 1.1 供应商地址、电话、联系人（R5C1 / R6C1 / R7C1） ──
    if not supplier_address and not supplier_phone and not supplier_contact:
        logger.warning(f"供应商'{supplier_name}'的地址/电话/联系人信息均缺失")
    _fill_cell(ws, 5, 1, f'供应商地址:{supplier_address}' if supplier_address else '供应商地址:')
    _fill_cell(ws, 6, 1, f'电话:{supplier_phone}' if supplier_phone else '电话:')
    _fill_cell(ws, 7, 1, f'联系人:{supplier_contact}' if supplier_contact else '联系人:')

    # ── 1.1 右上角日期（R2C36: 替换模板硬编码日期为当前日期） ──
    today = datetime.now()
    _fill_cell(ws, 2, 36, today, number_format='YYYY-M-D')

    # ── 1.2 打印预览设置：宽度缩放至1页，高度不限（动态适配） ──
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 999  # 设置足够大的值，等效于不限制纵向页数
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    # 清除固定的打印区域，让 Excel 自动检测数据范围
    ws.print_area = None

    # ── 2. 动态数据行 ──
    TEMPLATE_DATA_ROWS = 6   # 模板固定6个数据行 (R10-R15)
    DATA_START_ROW = 10      # 数据起始行
    DATA_ROW_HEIGHT = 20.25  # 数据行行高（与模板一致）

    # 每个数据行的9个合并列区域（与模板一致）
    DATA_COL_MERGES = [
        (1, 2),    # A:B   = 序号
        (3, 13),   # C:M   = 品名/规格/色号
        (14, 18),  # N:R   = 订购数量
        (19, 21),  # S:U   = 单位
        (22, 25),  # V:Y   = 单价/码
        (26, 29),  # Z:AC  = 税率
        (30, 33),  # AD:AG = 金额
        (34, 36),  # AH:AJ = 交货日期
        (37, 45),  # AK:AS = 备注
    ]

    total_rows = max(len(items), TEMPLATE_DATA_ROWS)  # 至少6行，多了就扩展

    # 如果物料超过6行，需要在 R16 位置插入额外行（openpyxl 自动下移后续行）
    extra_rows = total_rows - TEMPLATE_DATA_ROWS
    if extra_rows > 0:
        insert_pos = DATA_START_ROW + TEMPLATE_DATA_ROWS  # R16

        # 1) 先收集 footer 区域（R16-R22）的合并范围和单元格数据
        footer_start = insert_pos
        footer_merges = []
        for m in list(ws.merged_cells.ranges):
            if m.min_row >= footer_start:
                footer_merges.append((m.min_row, m.min_col, m.max_row, m.max_col))
                ws.unmerge_cells(start_row=m.min_row, start_column=m.min_col,
                                 end_row=m.max_row, end_column=m.max_col)

        footer_cells = {}
        for r in range(footer_start, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is not None:
                    footer_cells[(r, c)] = {
                        'value': cell.value,
                        'font': _copy_font_style(cell.font),
                        'alignment': _copy_align_style(cell.alignment),
                        'border': _copy_border_style(cell.border),
                        'number_format': cell.number_format,
                    }

        footer_heights = {}
        for r in range(footer_start, ws.max_row + 1):
            h = ws.row_dimensions[r].height
            if h:
                footer_heights[r] = h

        # 2) 清空 footer 区域的值
        for r in range(footer_start, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).value = None

        # 3) 插入空行（自动下移 footer 行，但合并范围已全部取消）
        ws.insert_rows(insert_pos, amount=extra_rows)

        # 4) 在新位置恢复 footer 合并和单元格数据
        for min_r, min_c, max_r, max_c in footer_merges:
            ws.merge_cells(start_row=min_r + extra_rows, start_column=min_c,
                           end_row=max_r + extra_rows, end_column=max_c)

        for (r, c), info in footer_cells.items():
            new_r = r + extra_rows
            cell = ws.cell(row=new_r, column=c)
            cell.value = info['value']
            cell.font = info['font']
            cell.alignment = info['alignment']
            cell.border = info['border']
            if info['number_format']:
                cell.number_format = info['number_format']

        for r, h in footer_heights.items():
            ws.row_dimensions[r + extra_rows].height = h

        # 5) 为新插入的行设置行高、全边框、合并区域和格式
        ref_row = DATA_START_ROW  # 以模板第一行数据行 R10 为格式参考
        ref_font = _copy_font_style(ws.cell(row=ref_row, column=1).font)
        ref_align = _copy_align_style(ws.cell(row=ref_row, column=1).alignment)
        from openpyxl.styles import Border as Bd, Side as Sd
        full_border = Bd(left=Sd(style='thin'), right=Sd(style='thin'),
                         top=Sd(style='thin'), bottom=Sd(style='thin'))

        for i in range(extra_rows):
            new_row = insert_pos + i
            ws.row_dimensions[new_row].height = DATA_ROW_HEIGHT

            # 先为整行所有列设置全边框（合并前设置，确保合并区域边缘有边框）
            for c in range(1, 46):
                cell = ws.cell(row=new_row, column=c)
                cell.border = full_border
                cell.font = ref_font
                cell.alignment = ref_align

            # 创建合并区域
            for c1, c2 in DATA_COL_MERGES:
                ws.merge_cells(start_row=new_row, start_column=c1, end_row=new_row, end_column=c2)

            # 合并起始单元格设置值（空白行）
            for c1, c2 in DATA_COL_MERGES:
                ws.cell(row=new_row, column=c1).value = None

    # ── 填充数据行 ──
    for idx in range(total_rows):
        row = DATA_START_ROW + idx
        if idx < len(items):
            item = items[idx]
            name_val = item['material_name']
            if item.get('specification'):
                name_val = f"{item['material_name']}（{item['specification']}）"
            qty = str(item['quantity'])

            _fill_cell(ws, row, 1, idx + 1)                        # 序号
            _fill_cell(ws, row, 3, name_val)                        # 品名/规格/色号
            _fill_cell(ws, row, 14, qty, number_format='@')         # 订购数量（文本格式）
            _fill_cell(ws, row, 19, item.get('unit') or '')         # 单位：BOM 数据
            _fill_cell(ws, row, 22, None)                            # 单价：留空
            _fill_cell(ws, row, 26, '含')                            # 税率：默认"含"
            _fill_cell(ws, row, 30, f'=N{row}*V{row}', number_format='#,##0.00')  # 金额=数量×单价
            _fill_cell(ws, row, 34, None)                            # 交货日期：留空
            _fill_cell(ws, row, 37, None)                            # 备注：留空
        else:
            # 空数据行：清空内容但保留格式
            _fill_cell(ws, row, 1, idx + 1)    # 序号保持编号
            _fill_cell(ws, row, 3, None)        # 品名清空
            _fill_cell(ws, row, 14, None)       # 数量清空
            _fill_cell(ws, row, 19, None)       # 单位清空
            _fill_cell(ws, row, 22, None)       # 单价清空
            _fill_cell(ws, row, 26, '含')        # 税率：默认"含"
            _fill_cell(ws, row, 30, f'=N{row}*V{row}', number_format='#,##0.00')  # 金额公式
            _fill_cell(ws, row, 34, None)       # 交货日期：留空
            _fill_cell(ws, row, 37, None)       # 备注清空

    # ── 3. 汇总金额（使用 Excel 公式自动计算） ──
    summary_row = DATA_START_ROW + total_rows  # 总金额行（自动下移后的位置）
    last_data_row = DATA_START_ROW + total_rows - 1  # 最后一个数据行

    # 总金额 = SUM(AD10:AD{last_data_row})
    _fill_cell(ws, summary_row, 30, f'=SUM(AD{DATA_START_ROW}:AD{last_data_row})', number_format='#,##0.00')
    # 税额行：税率输入字段（用户填入小数，如 0.13 表示 13%）
    _fill_cell(ws, summary_row + 1, 30, None, number_format='0%')
    # 应付金额 = 总金额 × (1 + 税率)
    _fill_cell(ws, summary_row + 2, 30, f'=AD{summary_row}*(1+AD{summary_row+1})', number_format='#,##0.00')

    return wb


@router.post("/orders/supplier-summary-export")
def export_supplier_summary_excel(
    data: OrderIdsRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    """按供应商维度导出采购订单汇总 Excel，单文件直接下载，多文件 ZIP 打包"""
    order_ids = data.ids

    try:
        sorted_suppliers = svc.query_supplier_groups(db, order_ids)
        if not sorted_suppliers:
            raise HTTPException(status_code=404, detail="未找到供应商数据")

        date_prefix = datetime.now().strftime('%Y-%m-%d')

        # 场景 A：仅1个供应商 → 直接导出 .xlsx
        if len(sorted_suppliers) == 1:
            supplier = sorted_suppliers[0]
            wb = create_supplier_workbook(supplier)
            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            wb.close()
            excel_buffer.seek(0)

            filename = f"{date_prefix}-{supplier['supplier_name']}_采购订单汇总.xlsx"
            encoded_filename = quote(filename, safe='')

            return StreamingResponse(
                excel_buffer,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
            )

        # 场景 B：2个及以上供应商 → ZIP 打包
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
            for supplier in sorted_suppliers:
                wb = create_supplier_workbook(supplier)
                excel_buffer = io.BytesIO()
                wb.save(excel_buffer)
                wb.close()
                excel_buffer.seek(0)
                filename = f"{supplier['supplier_name']}_采购订单汇总.xlsx"
                zf.writestr(filename, excel_buffer.read())

        zip_buffer.seek(0)
        zip_filename = f"{date_prefix}-供应商采购订单汇总.zip"
        encoded_filename = quote(zip_filename, safe='')

        return StreamingResponse(
            zip_buffer,
            media_type="application/zip",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.exception(f"导出供应商汇总Excel失败: {str(e)}")
        raise HTTPException(status_code=500, detail="导出失败，请稍后重试")


# ============================================================
# BOM 复制功能 (从 material.py 迁移)
# ============================================================


@router.get("/api/products/search")
def api_search_products_for_bom_copy(
    keyword: str = Query('', min_length=1),
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    """搜索产品用于复制BOM"""
    if not keyword:
        raise HTTPException(status_code=400, detail="请输入搜索关键词")

    products = svc.search_products_for_bom_copy(db, keyword)
    return {'success': True, 'products': products}


@router.get("/api/products/{source_product_id}/bom-preview")
def api_get_bom_preview(
    source_product_id: int,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    """获取源产品BOM预览用于复制确认"""
    try:
        result = svc.get_bom_preview(db, source_product_id)
        return {'success': True, **result}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/api/products/{target_product_id}/copy-bom")
def api_copy_bom_to_product(
    target_product_id: int,
    data: BomCopyRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    """复制BOM到目标产品"""
    try:
        result = svc.copy_bom_to_product(db, target_product_id, data.model_dump())
        return {'success': True, **result}
    except ValueError as e:
        error_msg = str(e)
        if '目标产品已有BOM' in error_msg:
            raise HTTPException(status_code=409, detail=error_msg, headers={'has_existing_bom': 'true'})
        raise HTTPException(status_code=400, detail=error_msg)


@router.post("/api/orders/batch-delete")
def api_batch_delete_purchase_orders(
    data: OrderIdsAltRequest,
    db: Session = Depends(get_db),
    _user=Depends(require_permission("product_manage")),
):
    """批量删除采购订单"""
    try:
        result = svc.batch_delete_purchase_orders(db, data.order_ids)
        return {'success': True, **result}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
