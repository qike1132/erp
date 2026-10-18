"""
采购中心 — Excel 导出工具函数

包含 HTML 打印模板生成和 Excel 工作簿创建逻辑。
"""

import os
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment as Al
from openpyxl.styles import Border as Bd
from openpyxl.styles import Font as Ft
from openpyxl.styles import Side as Sd

from fastapi_app.core.float_utils import round_qty
from fastapi_app.core.logger import logger

# 模板文件路径
_PROJECT_ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..', '..'))
_TEMPLATE_PATH = os.path.join(_PROJECT_ROOT, 'static', 'files', 'purchase_order_template.xlsx')

# 供应商采购订单模板的数据区布局（create_supplier_workbook 及其插行 helper 共用）
_TEMPLATE_DATA_ROWS = 6
_DATA_START_ROW = 10
_DATA_ROW_HEIGHT = 20.25
_DATA_COL_MERGES = [
    (1, 2), (3, 13), (14, 18), (19, 21),
    (22, 25), (26, 29), (30, 33), (34, 36), (37, 45),
]


def _validate_template_path(template_path: str) -> str:
    """验证模板文件路径在项目目录内，防止路径遍历"""
    normalized = os.path.normpath(template_path)
    if not normalized.startswith(_PROJECT_ROOT):
        raise ValueError("模板路径超出项目目录范围")
    if not os.path.isfile(normalized):
        raise FileNotFoundError(f"模板文件不存在: {normalized}")
    return normalized


def _copy_font_style(font):
    """复制 Font 对象"""
    return Ft(name=font.name, size=font.size, bold=font.bold, italic=font.italic,
              color=font.color, underline=font.underline, strike=font.strike)


def _copy_align_style(alignment):
    """复制 Alignment 对象"""
    return Al(horizontal=alignment.horizontal, vertical=alignment.vertical,
              wrap_text=alignment.wrap_text, text_rotation=alignment.text_rotation,
              indent=alignment.indent, shrink_to_fit=alignment.shrink_to_fit)


def _copy_border_style(border):
    """复制 Border 对象"""
    def _cp(side):
        return Sd(style=side.style, color=side.color) if side.style else Sd()
    return Bd(left=_cp(border.left), right=_cp(border.right),
              top=_cp(border.top), bottom=_cp(border.bottom))


def _fill_cell(ws, row, col, value, number_format=None):
    """向模板中的指定单元格填充值，保留原有格式"""
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if number_format:
        cell.number_format = number_format


def build_print_html(title: str, body_content: str, extra_css: str = "") -> str:
    """生成打印HTML页面的通用模板"""
    base_css = """body{font-family:"Microsoft YaHei",sans-serif;padding:20px}
.page-break{page-break-after:always}
.header{text-align:center;margin-bottom:20px;border-bottom:2px solid #333;padding-bottom:10px}
.header h1{margin:0;color:#333}
.print-date{text-align:right;color:#666;font-size:12px;margin-top:20px}
"""
    return f'''<!DOCTYPE html><html><head><meta charset="utf-8"><title>{title}</title>
<style>{base_css}{extra_css}</style></head><body>
{body_content}
<div class="print-date">打印时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</div></body></html>'''


def generate_normal_summary_print(orders, order_items_map):
    """生成普通格式打印HTML"""
    extra_css = """.order-section{margin-bottom:30px}
.order-title{font-size:16px;font-weight:bold;margin:15px 0 10px;padding:8px;background:#f5f5f5}
table{width:100%;border-collapse:collapse;margin-top:10px}
th,td{border:1px solid #ddd;padding:8px;text-align:left;font-size:12px}
th{background:#f5f5f5;font-weight:bold}
.order-meta{margin-bottom:10px;font-size:13px}"""
    body_parts = ['<div class="header"><h1>东莞市滨发玩具有限公司</h1><h2>采购订单汇总</h2></div>']
    for idx, order in enumerate(orders):
        if idx > 0:
            body_parts.append('<div class="page-break"></div>')
        items = order_items_map.get(order.id, [])
        body_parts.append(f'<div class="order-section"><div class="order-title">订单编号: {order.order_no}</div>')
        body_parts.append(f'<div class="order-meta"><strong>供应商:</strong> {order.supplier_name or "-"} | <strong>客户:</strong> {order.customer_name or "-"} | <strong>日期:</strong> {order.order_date or "-"}</div>')
        body_parts.append('<table><thead><tr><th>物料名称</th><th>规格</th><th>数量</th><th>单位</th><th>分类</th></tr></thead><tbody>')
        for item in items:
            body_parts.append(f'<tr><td>{item.material_name}</td><td>{item.specification or "-"}</td><td>{item.quantity}</td><td>{item.unit or "-"}</td><td>{item.material_category or "-"}</td></tr>')
        body_parts.append('</tbody></table></div>')
    return build_print_html("采购订单汇总", ''.join(body_parts), extra_css)


def generate_factory_summary_print(orders, order_items_map, source_orders=None):
    """生成工厂内部汇总格式打印HTML"""
    source_orders = source_orders or []
    source_order_map = {o.id: o for o in source_orders}

    grouped_data: dict[tuple[str, str], dict[str, dict[tuple[str, str, str], dict[str, Any]]]] = {}
    order_count_map = {}

    for order in orders:
        items = order_items_map.get(order.id, [])
        source_order = source_order_map.get(order.source_order_id) if order.source_order_id else None
        product_name = order.product_name
        if not product_name and source_order and source_order.product:
            product_name = source_order.product.name
        elif not product_name:
            product_name = order.customer_name or '未指定产品'
        product_qty = order.product_quantity
        if not product_qty and source_order:
            product_qty = source_order.quantity
        customer_name = source_order.customer if source_order else (order.customer_name or '未指定客户')
        cp_key = (customer_name, product_name)
        if cp_key not in grouped_data:
            grouped_data[cp_key] = {}
            order_count_map[cp_key] = {'order_count': 0, 'total_fabric_qty': 0, 'product_quantity': product_qty}
        for item in items:
            supplier_name = item.supplier_name or order.supplier_name or '未指定供应商'
            fabric_name = item.material_name or '未指定布料'
            fabric_spec = item.specification or ''
            fabric_qty = item.quantity or 0
            if supplier_name not in grouped_data[cp_key]:
                grouped_data[cp_key][supplier_name] = {}
            fabric_key = (fabric_name, fabric_spec, item.unit or '')
            if fabric_key not in grouped_data[cp_key][supplier_name]:
                grouped_data[cp_key][supplier_name][fabric_key] = {'fabric_name': fabric_name, 'specification': fabric_spec, 'quantity': fabric_qty, 'unit': item.unit or ''}
            else:
                existing = grouped_data[cp_key][supplier_name][fabric_key]
                existing['quantity'] = (float(existing['quantity']) if existing['quantity'] else 0) + (float(fabric_qty) if fabric_qty else 0)
            order_count_map[cp_key]['total_fabric_qty'] += fabric_qty

    for order in orders:
        source_order = source_order_map.get(order.source_order_id) if order.source_order_id else None
        product_name = order.product_name
        if not product_name and source_order and source_order.product:
            product_name = source_order.product.name
        elif not product_name:
            product_name = order.customer_name or '未指定产品'
        customer_name = source_order.customer if source_order else (order.customer_name or '未指定客户')
        cp_key = (customer_name, product_name)
        if cp_key in order_count_map:
            order_count_map[cp_key]['order_count'] += 1

    sorted_groups = {}
    for (customer_name, product_name), supplier_dict in grouped_data.items():
        suppliers_list: list[dict[str, Any]] = []
        for supplier_name, fabric_dict in supplier_dict.items():
            fabrics_list = list(fabric_dict.values())
            for f in fabrics_list:
                f['quantity'] = round_qty(float(f['quantity'])) if f['quantity'] else 0
            suppliers_list.append({'supplier_name': supplier_name, 'fabrics': fabrics_list})
        sorted_groups[(customer_name, product_name)] = suppliers_list

    extra_css = """@page{margin:0in}body{margin:0;padding:0}
.header h1{font-size:24pt;font-weight:bold}.header h2{margin:5px 0 0 0;font-size:18pt;font-weight:bold}
table{width:850.15pt;border-collapse:collapse;table-layout:fixed}
th,td{border:1.0pt solid #000000;padding:1px;text-align:center;vertical-align:middle;font-size:10.5pt;white-space:normal}
th{background:#E6E6E6;font-weight:bold}
.total-row{background-color:#F0F0F0;font-weight:bold}"""
    body_parts = ['<div class="header"><h1>东莞市滨发玩具有限公司</h1><h2>工厂内部采购订单汇总</h2></div>']
    body_parts.append('<table><thead><tr><th>客户</th><th>产品名称</th><th>订单数量</th><th>供应商</th><th>布料名称</th><th>规格型号</th><th>数量</th></tr></thead><tbody>')

    grand_total: float = 0
    for (customer_name, product_name), suppliers in sorted(sorted_groups.items()):
        total_fabric_rows = sum(len(s['fabrics']) for s in suppliers)
        order_info = order_count_map[(customer_name, product_name)]
        global_fabric_idx = 0
        for supplier in suppliers:
            for idx, fabric in enumerate(supplier['fabrics']):
                body_parts.append('<tr>')
                if global_fabric_idx == 0:
                    body_parts.append(f'<td rowspan="{total_fabric_rows}" style="font-weight:bold">{customer_name}</td>')
                    body_parts.append(f'<td rowspan="{total_fabric_rows}">{product_name}</td>')
                    body_parts.append(f'<td rowspan="{total_fabric_rows}">{float(order_info["product_quantity"]) if order_info["product_quantity"] else 0:g}</td>')
                if idx == 0:
                    body_parts.append(f'<td rowspan="{len(supplier["fabrics"])}">{supplier["supplier_name"]}</td>')
                body_parts.append(f'<td style="text-align:left">{fabric["fabric_name"]}</td>')
                body_parts.append(f'<td style="text-align:left">{fabric.get("specification", "")}</td>')
                body_parts.append(f'<td style="text-align:right">{fabric["quantity"]} {fabric["unit"]}</td>')
                body_parts.append('</tr>')
                grand_total += float(fabric["quantity"]) if fabric["quantity"] else 0
                global_fabric_idx += 1

    body_parts.append(f'<tfoot><tr class="total-row"><td colspan="4">合计</td><td></td><td></td><td style="text-align:right">{grand_total}</td></tr></tfoot>')
    body_parts.append('</table>')
    return build_print_html("工厂内部采购订单汇总", ''.join(body_parts), extra_css)


def generate_supplier_summary_print(supplier_groups, orders):
    """生成供应商汇总打印HTML"""
    extra_css = """@media print{@page{size:A4;margin:15mm}.no-print{display:none!important}}
.company-header h1{font-size:24px;font-weight:bold}.company-header h2{margin:10px 0 0 0;color:#666;font-size:18px}
.summary-info{margin:15px 0;padding:10px;background:#f9f9f9;border:1px solid #ddd}
.supplier-section{margin-bottom:30px;border:1px solid #ccc}
.supplier-header{background:#4a90e2;color:white;padding:12px 15px;font-size:16px;font-weight:bold}
.supplier-info{padding:10px 15px;background:#f5f5f5;border-bottom:1px solid #ddd}
table{width:100%;border-collapse:collapse}
th,td{border:1px solid #ddd;padding:8px;text-align:left;vertical-align:middle}
th{background:#e8e8e8;font-weight:bold;text-align:center}
.total-row{background:#fff3cd;font-weight:bold}
.footer{margin-top:30px;text-align:right;color:#666;font-size:11px}"""
    body_parts = ['<div class="company-header"><h1>东莞市滨发玩具有限公司</h1><h2>供应商采购订单汇总</h2></div>']
    body_parts.append(f'<div class="summary-info"><strong>汇总时间：</strong>{datetime.now().strftime("%Y-%m-%d %H:%M:%S")} | <strong>供应商数量：</strong>{len(supplier_groups)} | <strong>订单数量：</strong>{len(orders)}</div>')
    for idx, supplier in enumerate(supplier_groups):
        if idx > 0:
            body_parts.append('<div class="page-break"></div>')
        total_qty = sum(float(item['quantity']) for item in supplier['items'])
        body_parts.append(f'<div class="supplier-section"><div class="supplier-header">供应商：{supplier["supplier_name"]}</div>')
        body_parts.append(f'<div class="supplier-info"><strong>关联订单：</strong>{", ".join(supplier["orders"])} | <strong>物料种类：</strong>{len(supplier["items"])} 种</div>')
        body_parts.append('<table><thead><tr><th>序号</th><th>商品名称</th><th>规格型号</th><th>单位</th><th>数量</th><th>物料分类</th><th>关联订单</th></tr></thead><tbody>')
        for item_idx, item in enumerate(supplier['items'], 1):
            body_parts.append(f'<tr><td style="text-align:center">{item_idx}</td><td>{item["material_name"]}</td><td>{item["specification"] or "-"}</td><td style="text-align:center">{item["unit"] or "-"}</td><td style="text-align:right">{float(item["quantity"]):,.2f}</td><td style="text-align:center">{item["material_category"] or "-"}</td><td>共{item["order_count"]}单</td></tr>')
        body_parts.append(f'<tfoot><tr class="total-row"><td colspan="4" style="text-align:right">合计：</td><td style="text-align:right">{total_qty:,.2f}</td><td colspan="2"></td></tr></tfoot></table></div>')
    return build_print_html("供应商采购订单汇总", ''.join(body_parts), extra_css)


def _qty_display_text(value) -> str:
    """数量展示文本：整数不带小数位、非整数最多 2 位小数。

    与前端 ``qtyDisplayText`` 同口径，保证导出 Excel 的「数量」与页面逐字一致
    （旧实现直接写 ``str(float)``，整数会渲染成 "20.0"）。
    """
    rounded = round_qty(float(value or 0))
    if rounded == int(rounded):
        return str(int(rounded))
    return f'{rounded:.2f}'.rstrip('0').rstrip('.')


def _insert_supplier_data_rows(ws, extra_rows: int) -> None:
    """在模板数据区下方插入 extra_rows 行，并保持页脚（合计/签名区）的样式与合并。

    模板固定 6 行数据，供应商物料超出时必须插行；插行会顶掉下方的页脚，
    因此先把页脚的合并区域、单元格内容/样式、行高整体 snapshot，清空后插行，
    再按新行号原样还原。抽成独立函数是为了控制 create_supplier_workbook 长度
    （2026-10-18：该函数 125 行，超出 function_length 棘轮基线 1）。
    """
    DATA_START_ROW = _DATA_START_ROW
    TEMPLATE_DATA_ROWS = _TEMPLATE_DATA_ROWS
    insert_pos = DATA_START_ROW + TEMPLATE_DATA_ROWS
    footer_start = insert_pos

    footer_merges = []
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= footer_start:
            footer_merges.append((m.min_row, m.min_col, m.max_row, m.max_col))
            ws.unmerge_cells(start_row=m.min_row, start_column=m.min_col,
                             end_row=m.max_row, end_column=m.max_col)

    footer_cells = {}
    for r in range(footer_start, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is not None:
                footer_cells[(r, c)] = {
                    'value': cell.value,
                    'font': _copy_font_style(cell.font),
                    'alignment': _copy_align_style(cell.alignment),
                    'border': _copy_border_style(cell.border),
                    'number_format': cell.number_format,
                }

    footer_heights = {}
    for r in range(footer_start, ws.max_row + 1):
        h = ws.row_dimensions[r].height
        if h:
            footer_heights[r] = h

    for r in range(footer_start, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None

    ws.insert_rows(insert_pos, amount=extra_rows)

    for min_r, min_c, max_r, max_c in footer_merges:
        ws.merge_cells(start_row=min_r + extra_rows, start_column=min_c,
                       end_row=max_r + extra_rows, end_column=max_c)

    for (r, c), info in footer_cells.items():
        new_r = r + extra_rows
        cell = ws.cell(row=new_r, column=c)
        cell.value = info['value']
        cell.font = info['font']
        cell.alignment = info['alignment']
        cell.border = info['border']
        if info['number_format']:
            cell.number_format = info['number_format']

    for r, h in footer_heights.items():
        ws.row_dimensions[r + extra_rows].height = h

    ref_row = _DATA_START_ROW
    ref_font = _copy_font_style(ws.cell(row=ref_row, column=1).font)
    ref_align = _copy_align_style(ws.cell(row=ref_row, column=1).alignment)
    full_border = Bd(left=Sd(style='thin'), right=Sd(style='thin'),
                     top=Sd(style='thin'), bottom=Sd(style='thin'))

    for i in range(extra_rows):
        new_row = insert_pos + i
        ws.row_dimensions[new_row].height = _DATA_ROW_HEIGHT

        for c in range(1, 46):
            cell = ws.cell(row=new_row, column=c)
            cell.border = full_border
            cell.font = ref_font
            cell.alignment = ref_align

        for c1, c2 in _DATA_COL_MERGES:
            ws.merge_cells(start_row=new_row, start_column=c1, end_row=new_row, end_column=c2)

        for c1, _ in _DATA_COL_MERGES:
            ws.cell(row=new_row, column=c1).value = None


def create_supplier_workbook(supplier: dict) -> Workbook:
    """基于模板文件为单个供应商创建采购订单 Excel 工作簿"""
    template_path = os.path.normpath(_TEMPLATE_PATH)
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    supplier_name = supplier['supplier_name']
    items = supplier['items']
    supplier_address = supplier.get('address', '') or ''
    supplier_phone = supplier.get('phone', '') or ''
    supplier_contact = supplier.get('contact_person', '') or ''

    _fill_cell(ws, 4, 1, f'供应商名称:{supplier_name}')
    if not supplier_address and not supplier_phone and not supplier_contact:
        logger.warning(f"供应商'{supplier_name}'的地址/电话/联系人信息均缺失")
    _fill_cell(ws, 5, 1, f'供应商地址:{supplier_address}' if supplier_address else '供应商地址:')
    _fill_cell(ws, 6, 1, f'电话:{supplier_phone}' if supplier_phone else '电话:')
    _fill_cell(ws, 7, 1, f'联系人:{supplier_contact}' if supplier_contact else '联系人:')

    today = datetime.now()
    _fill_cell(ws, 2, 36, today, number_format='YYYY-M-D')

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 999
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_area = None

    # 数据区布局用模块级常量（_insert_supplier_data_rows 共用）；此处起本地别名
    # 只为让下方填充逻辑保持可读。
    TEMPLATE_DATA_ROWS = _TEMPLATE_DATA_ROWS
    DATA_START_ROW = _DATA_START_ROW

    total_rows = max(len(items), TEMPLATE_DATA_ROWS)
    extra_rows = total_rows - TEMPLATE_DATA_ROWS

    if extra_rows > 0:
        _insert_supplier_data_rows(ws, extra_rows)

    for idx in range(total_rows):
        row = DATA_START_ROW + idx
        if idx < len(items):
            item = items[idx]
            name_val = item['material_name']
            if item.get('specification'):
                name_val = f"{item['material_name']}（{item['specification']}）"
            qty = _qty_display_text(item['quantity'])

            _fill_cell(ws, row, 1, idx + 1)
            _fill_cell(ws, row, 3, name_val)
            _fill_cell(ws, row, 14, qty, number_format='@')
            _fill_cell(ws, row, 19, item.get('unit') or '')
            _fill_cell(ws, row, 22, None)
            _fill_cell(ws, row, 26, '含')
            _fill_cell(ws, row, 30, f'=N{row}*V{row}', number_format='#,##0.00')
            _fill_cell(ws, row, 34, None)
            _fill_cell(ws, row, 37, None)
        else:
            _fill_cell(ws, row, 1, idx + 1)
            _fill_cell(ws, row, 3, None)
            _fill_cell(ws, row, 14, None)
            _fill_cell(ws, row, 19, None)
            _fill_cell(ws, row, 22, None)
            _fill_cell(ws, row, 26, '含')
            _fill_cell(ws, row, 30, f'=N{row}*V{row}', number_format='#,##0.00')
            _fill_cell(ws, row, 34, None)
            _fill_cell(ws, row, 37, None)

    summary_row = DATA_START_ROW + total_rows
    last_data_row = DATA_START_ROW + total_rows - 1

    _fill_cell(ws, summary_row, 30, f'=SUM(AD{DATA_START_ROW}:AD{last_data_row})', number_format='#,##0.00')
    _fill_cell(ws, summary_row + 1, 30, None, number_format='0%')
    _fill_cell(ws, summary_row + 2, 30, f'=AD{summary_row}*(1+AD{summary_row+1})', number_format='#,##0.00')

    return wb
