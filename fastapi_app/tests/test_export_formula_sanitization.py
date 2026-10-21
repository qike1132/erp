"""审查 H12：生产/成品库存导出的 Excel 公式注入清洗回归测试。

此前 `_sanitize_excel_value` 只接了采购打印与 report_exporter；生产侧三处
导出（production_summary 手工部/包装部、daily_production 生产报表、
finished_stock 端点）把 worker_name/remark/customer/product_name 等**用户
可控文本**直入单元格——Excel 打开即按公式/DDE 解析。

断言打在「单元格实际落到 Excel 的那串字符」上（load_workbook 读回），
与 test_purchase_export_sanitize.py 同口径：清洗后的值以 `\\t` 前缀开头
（打断公式解析的防御标记），不得以 `=`/`+`/`-`/`@` 开头。
"""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import openpyxl
import pytest

_FORMULA_LEADERS = ('=', '+', '-', '@')


def _assert_no_formula_leader(xlsx_bytes: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    bad: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v[:1] in _FORMULA_LEADERS:
                    bad.append(f'{ws.title}!{cell.coordinate}={v!r}')
    assert bad == [], f'存在以公式引导符开头的单元格（Excel 打开即解析）: {bad}'
    return bad


def _sanitized_cells(xlsx_bytes: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    return [
        c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith('\t')
    ]


class _FakeScalars:
    def __init__(self, rows):
        self._rows = rows

    def all(self):
        return self._rows


class _FakeResult:
    def __init__(self, rows):
        self._rows = rows

    def scalars(self):
        return _FakeScalars(self._rows)


class _FakeDb:
    """只实现被测导出用到的最小面：scalars() / execute()。"""

    def __init__(self, rows):
        self._rows = rows

    async def scalars(self, _query):
        return _FakeScalars(self._rows)

    async def execute(self, _query):
        return _FakeResult(self._rows)


# ---------------------------------------------------------------------------
# ① 成品库存端点（纯同步，直接行为断言）
# ---------------------------------------------------------------------------

def test_finished_stock_excel_sanitizes_formula_cells():
    from fastapi_app.api.v1.endpoints.finished_stock import _build_finished_stock_excel

    rows = [{
        'customer_name': '=cmd|\' /C calc\'!A0',
        'product_name': '=1+1',
        'barcode': '@SUM(A1)',
        'batch_no': '-2+3',
        'batch_available': 5,
        'is_custom': False,
        'custom_country': '+86',
    }]

    xlsx = _build_finished_stock_excel(rows)

    _assert_no_formula_leader(xlsx)
    sanitized = _sanitized_cells(xlsx)
    assert len(sanitized) == 5  # 五个文本字段全部走了清洗器
    assert sanitized[0] == '\t=cmd|\' /C calc\'!A0'


# ---------------------------------------------------------------------------
# ② 生产汇总（手工部 / 包装部）
# ---------------------------------------------------------------------------

async def test_craft_works_export_sanitizes_worker_and_remark():
    from fastapi_app.services import production_summary_service as pss

    records = [
        SimpleNamespace(
            work_date=date(2026, 10, 21), worker_name='=cmd|\' /C calc\'!A0',
            quantity=5, remark='@SUM(A1)', create_time=None,
        ),
    ]

    xlsx = await pss.export_craft_works(_FakeDb(records), '', '')

    _assert_no_formula_leader(xlsx)
    assert '\t=cmd|\' /C calc\'!A0' in _sanitized_cells(xlsx)


async def test_packaging_export_sanitizes_customer_and_product():
    from fastapi_app.services import production_summary_service as pss

    records = [
        SimpleNamespace(
            pack_date=date(2026, 10, 21), customer='=1+1', product_name='@SUM(A1)',
            barcode='-2+3', quantity=3, remark='=x', create_time=None,
        ),
    ]

    xlsx = await pss.export_packaging_records(_FakeDb(records), '', '')

    _assert_no_formula_leader(xlsx)
    # 客户/产品/69码/备注 四个文本字段都应带清洗前缀
    assert len(_sanitized_cells(xlsx)) == 4


# ---------------------------------------------------------------------------
# ③ 生产报表（daily_production）
# ---------------------------------------------------------------------------

async def test_daily_production_export_sanitizes_user_fields(monkeypatch):
    from fastapi_app.services import daily_production_service as dps
    from fastapi_app.services import export_limits

    monkeypatch.setattr(
        export_limits, 'ensure_within_export_limit', AsyncMock(), raising=False)
    monkeypatch.setattr(
        dps, '_build_product_price_fallback', AsyncMock(return_value={}))

    records = [
        SimpleNamespace(
            production_date=date(2026, 10, 21), department='filling',
            product_name='=1+1', barcode='@SUM(A1)', unit_price=0,
            quantity=2, remark='=cmd|\' /C calc\'!A0', created_by='-2+3',
            create_time=None, product_id=None,
        ),
    ]

    xlsx = await dps.export_records(
        _FakeDb(records), start_date='2026-10-01', end_date='2026-10-21')

    _assert_no_formula_leader(xlsx)
    # 产品名称/69码/备注/创建人 四个文本字段都应带清洗前缀
    assert len(_sanitized_cells(xlsx)) == 4


# ---------------------------------------------------------------------------
# ④ 护栏回归（H13）：dry-run 不放拦截、真写必须先过授权口
# ---------------------------------------------------------------------------

def test_cleanup_script_guard_respects_dry_run():
    """--apply 才拦；默认 dry-run（只读）不得被护栏挡住。"""
    root = Path(__file__).resolve().parent.parent.parent
    text = (root / 'scripts' / 'cleanup_duplicate_purchase_orders.py').read_text(encoding='utf-8')
    assert 'if args.apply:' in text
    assert text.index('if args.apply:') < text.index('require_prod_write_consent')


def test_backfill_script_guard_respects_execute_flag():
    root = Path(__file__).resolve().parent.parent.parent
    text = (root / 'scripts' / 'backfill_production_stock_20261006.py').read_text(encoding='utf-8')
    assert 'if args.execute:' in text
    assert text.index('if args.execute:') < text.index('require_prod_write_consent')


@pytest.mark.parametrize('rel,flag', [
    ('scripts/migrate_password_hashes.py', 'if not args.dry_run:'),
    ('scripts/perf_add_indexes.py', None),
    ('scripts/reset_admin_password.py', None),
])
def test_guard_wired_in_scripts(rel, flag):
    """H13：五个脚本必须接线统一护栏（无条件或条件守卫）。"""
    from pathlib import Path
    root = Path(__file__).resolve().parent.parent.parent
    text = (root / rel).read_text(encoding='utf-8')
    assert 'require_prod_write_consent' in text
    if flag:
        assert flag in text
