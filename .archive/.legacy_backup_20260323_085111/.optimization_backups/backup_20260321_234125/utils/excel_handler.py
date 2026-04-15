"""
统一的Excel文件处理工具
用于批量导入等场景,减少重复代码
"""
import tempfile
import os
from datetime import datetime, timedelta
from contextlib import contextmanager


@contextmanager
def excel_reader(file):
    """
    Excel文件读取上下文管理器
    自动处理临时文件的创建和清理
    
    参数:
        file: Flask上传的文件对象
    
    用法:
        with excel_reader(file) as reader:
            if reader.error:
                return jsonify({'success': False, 'error': reader.error})
            
            for row in range(2, reader.max_row + 1):
                value = reader.get_cell_value(row, col)
    
    返回:
        ExcelReader对象
    """
    reader = ExcelReader(file)
    try:
        yield reader
    finally:
        reader.cleanup()


class ExcelReader:
    """Excel文件读取器"""
    
    def __init__(self, file):
        self.file = file
        self.temp_path = None
        self.temp_fd = None
        self.wb = None
        self.ws = None
        self.error = None
        self.max_row = 0
        self.max_col = 0
        
        self._initialize()
    
    def _initialize(self):
        """初始化Excel读取器"""
        # 检查文件
        if not self.file or self.file.filename == '':
            self.error = '没有选择文件'
            return
        
        # 检查文件扩展名
        filename_lower = self.file.filename.lower()
        if not (filename_lower.endswith('.xlsx') or filename_lower.endswith('.xls')):
            self.error = '只支持 Excel 文件格式 (.xlsx 或 .xls)'
            return
        
        # 创建临时文件
        suffix = '.xlsx' if filename_lower.endswith('.xlsx') else '.xls'
        self.temp_fd, self.temp_path = tempfile.mkstemp(suffix=suffix)
        
        try:
            # 保存文件
            self.file.save(self.temp_path)
            
            # 根据文件类型打开
            if filename_lower.endswith('.xlsx'):
                self._open_xlsx()
            else:
                self._open_xls()
                
        except Exception as e:
            self.error = self._handle_excel_error(e)
    
    def _open_xlsx(self):
        """打开xlsx格式文件"""
        try:
            from openpyxl import load_workbook
            
            # 验证文件格式
            try:
                wb_test = load_workbook(filename=self.temp_path, read_only=True)
                wb_test.close()
            except Exception as format_error:
                self.error = self._handle_excel_error(format_error)
                return
            
            # 正式打开文件
            self.wb = load_workbook(self.temp_path)
            self.ws = self.wb.active
            self.max_row = self.ws.max_row
            self.max_col = self.ws.max_column
            
        except ImportError:
            self.error = '系统缺少 openpyxl 库,请安装: pip install openpyxl'
    
    def _open_xls(self):
        """打开xls格式文件"""
        try:
            import xlrd
            
            self.wb = xlrd.open_workbook(self.temp_path)
            self.ws = self.wb.sheet_by_index(0)
            self.max_row = self.ws.nrows
            self.max_col = self.ws.ncols
            
        except ImportError:
            self.error = '系统缺少 xlrd 库,请安装: pip install xlrd==1.2.0'
        except Exception as e:
            self.error = self._handle_excel_error(e)
    
    def get_cell_value(self, row, col):
        """
        获取单元格值
        
        参数:
            row: 行号(从1开始)
            col: 列号(从1开始)
        
        返回:
            单元格值
        """
        if self.error:
            return None
        
        filename_lower = self.file.filename.lower()
        
        if filename_lower.endswith('.xlsx'):
            # openpyxl方式
            return self.ws.cell(row=row, column=col).value
        else:
            # xlrd方式(行列从0开始)
            try:
                val = self.ws.cell_value(row - 1, col - 1)
                # 处理数字类型
                if isinstance(val, float) and val == int(val):
                    return int(val)
                return val
            except:
                return None
    
    def validate_headers(self, expected_headers, start_col=1):
        """
        验证表头
        
        参数:
            expected_headers: 期望的表头列表
            start_col: 起始列(默认1)
        
        返回:
            (success, actual_headers, error_message)
        """
        if self.error:
            return False, [], self.error
        
        # 读取实际表头
        actual_headers = []
        for col in range(start_col, start_col + len(expected_headers)):
            cell_value = self.get_cell_value(row=1, col=col)
            if cell_value is None:
                actual_headers.append('')
            else:
                actual_headers.append(str(cell_value).strip())
        
        # 检查表头数量
        if len(actual_headers) < len(expected_headers):
            return False, actual_headers, f'表头列数不足,期望{len(expected_headers)}列'
        
        # 验证表头内容
        for i, expected in enumerate(expected_headers):
            if actual_headers[i].replace(' ', '') != expected.replace(' ', ''):
                return False, actual_headers, f'表头格式不正确,期望:{expected_headers},实际:{actual_headers}'
        
        return True, actual_headers, None
    
    def check_data_rows(self, min_rows=1):
        """
        检查是否有数据行
        
        参数:
            min_rows: 最少数据行数
        
        返回:
            (success, error_message)
        """
        if self.error:
            return False, self.error
        
        if self.max_row < 1 + min_rows:  # 表头 + 数据行
            return False, f'Excel文件中没有数据行,请至少填写{min_rows}行数据'
        
        return True, None
    
    def parse_date_value(self, date_value):
        """
        解析日期值
        
        参数:
            date_value: 单元格中的日期值
        
        返回:
            datetime对象
        """
        if date_value is None or str(date_value).strip() == '':
            return datetime.now()
        
        try:
            # 如果是datetime对象
            if isinstance(date_value, datetime):
                return date_value
            
            # 如果是Excel日期(数字格式)
            if isinstance(date_value, (int, float)):
                base_date = datetime(1899, 12, 30)
                return base_date + timedelta(days=int(date_value))
            
            # 如果是字符串
            if isinstance(date_value, str):
                date_str = date_value.strip()
                # 尝试多种日期格式
                date_formats = [
                    '%Y-%m-%d', 
                    '%Y/%m/%d', 
                    '%Y-%m-%d %H:%M:%S', 
                    '%Y/%m/%d %H:%M:%S'
                ]
                for fmt in date_formats:
                    try:
                        return datetime.strptime(date_str, fmt)
                    except ValueError:
                        continue
            
            return datetime.now()
        except Exception:
            return datetime.now()
    
    def cleanup(self):
        """清理临时文件"""
        try:
            if self.temp_fd:
                os.close(self.temp_fd)
            if self.temp_path and os.path.exists(self.temp_path):
                os.unlink(self.temp_path)
        except:
            pass
    
    @staticmethod
    def _handle_excel_error(error):
        """处理Excel错误"""
        error_msg = str(error).lower()
        
        if 'no valid workbook' in error_msg or 'zip' in error_msg:
            return 'Excel文件格式无效或已损坏,请使用Microsoft Excel重新保存为标准.xlsx格式'
        elif 'unsupported format' in error_msg or 'corrupt' in error_msg:
            return 'Excel文件格式无效或已损坏,请使用Microsoft Excel重新保存'
        else:
            return f'Excel文件错误: {str(error)}'


def validate_required_field(value, field_name, row_num=None):
    """
    验证必填字段
    
    参数:
        value: 字段值
        field_name: 字段名称
        row_num: 行号(可选)
    
    返回:
        (success, error_message)
    """
    if value is None or str(value).strip() == '':
        if row_num:
            return False, f'第{row_num}行:{field_name}不能为空'
        else:
            return False, f'{field_name}不能为空'
    return True, None


def parse_quantity(quantity_value, field_name='数量', row_num=None):
    """
    解析数量字段
    
    参数:
        quantity_value: 数量值
        field_name: 字段名称
        row_num: 行号(可选)
    
    返回:
        (quantity, error_message)
    """
    if quantity_value is None:
        if row_num:
            return None, f'第{row_num}行:{field_name}不能为空'
        else:
            return None, f'{field_name}不能为空'
    
    try:
        # 尝试转换为整数
        if isinstance(quantity_value, str):
            quantity_value = quantity_value.strip()
            if quantity_value == '':
                if row_num:
                    return None, f'第{row_num}行:{field_name}不能为空'
                else:
                    return None, f'{field_name}不能为空'
            quantity = int(float(quantity_value))
        elif isinstance(quantity_value, float):
            quantity = int(quantity_value)
        else:
            quantity = int(quantity_value)
        
        if quantity <= 0:
            if row_num:
                return None, f'第{row_num}行:{field_name}必须大于0'
            else:
                return None, f'{field_name}必须大于0'
        
        return quantity, None
    except (ValueError, TypeError):
        if row_num:
            return None, f'第{row_num}行:{field_name}必须是有效的数字,实际值:{quantity_value}'
        else:
            return None, f'{field_name}必须是有效的数字'
