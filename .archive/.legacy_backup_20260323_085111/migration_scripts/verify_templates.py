#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
验证 Excel 模板文件内容
"""

from openpyxl import load_workbook
import os

def verify_template(filepath):
    """验证模板文件内容"""
    print(f"\n{'='*60}")
    print(f"验证文件：{os.path.basename(filepath)}")
    print(f"{'='*60}")
    
    wb = load_workbook(filepath)
    ws = wb.active
    
    print(f"工作表名称：{ws.title}")
    print(f"\n列头信息：")
    
    # 读取列头
    headers = []
    for col in range(1, 7):
        cell = ws.cell(row=1, column=col)
        headers.append(cell.value)
        print(f"  列{col}: {cell.value}")
    
    print(f"\n示例数据（共{ws.max_row - 1}行）：")
    
    # 读取示例数据
    for row in range(2, ws.max_row + 1):
        row_data = []
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            row_data.append(str(cell.value) if cell.value else '')
        print(f"  行{row}: {' | '.join(row_data)}")
    
    print(f"\n样式信息：")
    header_cell = ws.cell(row=1, column=1)
    print(f"  列头字体加粗：{header_cell.font.bold}")
    print(f"  列头对齐方式：{header_cell.alignment.horizontal}")
    
    print(f"\n列宽信息：")
    for col_idx in range(1, 7):
        col_letter = chr(64 + col_idx)
        width = ws.column_dimensions[col_letter].width
        print(f"  列{col_letter}: {width}")
    
    return True


def main():
    """主函数"""
    templates = [
        'static/files/订单批量上传模板.xlsx',
        'static/files/入库单批量上传模板.xlsx',
        'static/files/出库单批量上传模板.xlsx'
    ]
    
    print("开始验证 Excel 模板文件...\n")
    
    for template in templates:
        if os.path.exists(template):
            verify_template(template)
        else:
            print(f"✗ 文件不存在：{template}")
    
    print(f"\n{'='*60}")
    print("✓ 所有模板文件验证完成！")
    print(f"{'='*60}")


if __name__ == '__main__':
    main()
