"""ReportExporter 导出冒烟测试

背景：五类报表导出（订单/备货/出货/抵扣/退货）此前全仓无单测（2026-10-06 迭代审核 P2-2）。
本文件参数化覆盖：
- 每类导出的工作表名、表头、行值映射
- Excel 公式注入清洗（=+-@ 开头字符串加制表符前缀）
- generate_filename 的中文文件名映射
"""
import io

import pytest
from openpyxl import load_workbook

from utils.report_exporter import ReportExporter, generate_filename

# (导出方法, 工作表名, 表头, 行值取值键)
# 2026-10-13：表头补「单号/订单号」与「条码」两列（把行里已有的字段渲染出来）。
_CASES = [
    (
        'export_order_logs', '订单流水',
        ['下单时间', '订单号', '客户', '产品名称', '条码', '数量', '备注'],
        ['order_date', 'order_no', 'customer', 'product_name', 'product_barcode', 'quantity', 'note'],
    ),
    (
        'export_stockin_logs', '备货流水',
        ['备货时间', '备货单号', '客户', '产品名称', '条码', '备货数量', '备注'],
        ['stockin_date', 'record_no', 'customer', 'product_name', 'product_barcode', 'quantity', 'note'],
    ),
    (
        'export_stockout_logs', '出货流水',
        ['出货时间', '出货单号', '客户', '产品名称', '条码', '出货数量', '备注'],
        ['stockout_date', 'record_no', 'customer', 'product_name', 'product_barcode', 'quantity', 'note'],
    ),
    (
        'export_deduction_logs', '备货抵扣流水',
        ['抵扣时间', '出货单号', '订单号', '客户', '产品名称', '条码', '抵扣数量', '订单备注'],
        ['deduction_time', 'record_no', 'order_no', 'customer', 'product_name', 'product_barcode', 'quantity', 'order_note'],
    ),
    (
        'export_return_logs', '退货流水',
        ['退货时间', '退货单号', '客户', '产品名称', '条码', '退货数量', '备注'],
        ['return_date', 'record_no', 'customer', 'product_name', 'product_barcode', 'quantity', 'note'],
    ),
]

# 各类行键的全集（构造日志行时缺省补空，模拟真实行字典）
_ALL_KEYS = sorted({k for _, _, _, keys in _CASES for k in keys})


def _log(note: str, note_key: str = 'note') -> dict:
    row = dict.fromkeys(_ALL_KEYS, '')
    row.update({'customer': '测试客户', 'product_name': '测试产品', 'quantity': 12, note_key: note})
    return row


def _export(method: str, logs: list[dict]):
    data: bytes = getattr(ReportExporter(), method)(logs, {'total_count': len(logs)})
    return load_workbook(io.BytesIO(data)).active


def _cell(ws, row: int, col: int):
    """openpyxl 把空串写为空单元格，读回是 None；统一归一成 '' 再断言"""
    return ws.cell(row=row, column=col).value or ''


class TestExportSmoke:
    """五类导出的表名/表头/行值/公式注入清洗"""

    @pytest.mark.parametrize('method,sheet,headers,keys', _CASES)
    def test_sheet_headers_and_rows(
        self, method: str, sheet: str, headers: list[str], keys: list[str],
    ):
        log = _log('正常备注', keys[-1])
        ws = _export(method, [log])

        assert ws.title == sheet
        assert [c.value for c in ws[1]] == headers
        assert [_cell(ws, 2, i + 1) for i in range(len(keys))] == [log[k] for k in keys]

    @pytest.mark.parametrize('method,keys', [(m, k) for m, _, _, k in _CASES])
    def test_formula_injection_sanitized(self, method: str, keys: list[str]):
        """'=' 开头的备注被加制表符前缀，阻止 Excel 公式解析"""
        note_key = keys[-1]
        ws = _export(method, [_log('=1+1', note_key)])

        assert _cell(ws, 2, len(keys)) == '\t=1+1'

    def test_export_bytes_are_xlsx(self):
        """返回值是真实的 xlsx 字节（zip 魔数，可被 openpyxl 读回）"""
        ws = _export('export_order_logs', [_log('x')])
        assert ws.max_row >= 2

    def test_generate_filename(self):
        assert generate_filename('return') == '退货流水.xlsx'
        assert generate_filename('order') == '订单流水.xlsx'
        assert generate_filename('nope') == '报表.xlsx'


def _inventory_row(**over) -> dict:
    """构造一行库存汇总数据（字段名对齐 stock_inventory_service._format_inventory_rows）"""
    row = {
        'customer': '测试客户',
        'product_name': '测试产品',
        'order_quantity': 12,
        'stockin_quantity': 5,
        'total_stock': 17,
        'inventory_balance': 8,
    }
    row.update(over)
    return row


def _row_values(ws, row: int, ncols: int) -> list:
    """读回整行原始值（不把 0 归一成空串——数值 0 / 负数要能断言）"""
    return [ws.cell(row=row, column=c).value for c in range(1, ncols + 1)]


_INVENTORY_HEADERS = ['客户', '产品名称', '订单数量', '备货数量', '总订单', '库存']


class TestInventoryExport:
    """库存汇总导出：表头与行须与 /stock 页表格 6 列同口径。

    历史缺陷（2026-10-10）：导出只有 5 列，第 5 列表头误写成「总库存」
    （应为「总订单」），且整列漏掉了「库存」。本测试锁定正确口径。
    """

    def test_headers_and_row_match_stock_page_columns(self):
        data = ReportExporter().export_inventory_data([_inventory_row()], {'total_count': 1})
        ws = load_workbook(io.BytesIO(data)).active

        assert ws.title == '库存汇总'
        assert [c.value for c in ws[1]] == _INVENTORY_HEADERS
        assert _row_values(ws, 2, 6) == ['测试客户', '测试产品', 12, 5, 17, 8]

    def test_inventory_balance_keeps_zero_and_negative(self):
        """库存列 0 与负数原样输出（表格口径），不能变成空值"""
        data = ReportExporter().export_inventory_data(
            [_inventory_row(inventory_balance=0),
             _inventory_row(product_name='负结余', inventory_balance=-2500)],
            {'total_count': 2})
        ws = load_workbook(io.BytesIO(data)).active

        assert ws.cell(row=2, column=6).value == 0
        assert ws.cell(row=3, column=6).value == -2500
