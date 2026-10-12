"""产品导入表头改版（2026-10-12 用户裁定）—— 客户 / 工厂 / 价格的识别与回填

背景：模板「产品管理」Sheet 改名为「新品录入」，表头由
``客户名称 | 产品名称 | 69码 | 产品价格 | 图片``
改为 ``客户 | 产品名称 | 工厂 | 69码 | 产品价格 | 图片``；「备货管理」表头统一为
「备货数量」。用户此前反馈「批量上传后没有客户和价格」，根因有二：
① 产品导入从不读「客户名称」列（客户归属永远不写）；
② 已存在的产品整行跳过、只覆盖图片，价格填了也进不去，界面还提示「导入成功」。

本文件同时锁住模板资产本身（表头 + 单元格图片）与两条导入实现（端点用的
``import_products_service``、综合导入/ARQ 用的 ``import_products_from_data``）。
"""
import io
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select

from fastapi_app.services import batch_import_service, storage_service
from fastapi_app.shared_models import Customer, Product, ProductCustomer

TEMPLATE = Path(__file__).resolve().parent.parent / 'static' / 'files' / '统一导入模板.xlsx'

HEADERS = ['客户', '产品名称', '工厂', '69码', '产品价格', '图片']


class _Upload:
    """满足 ``schemas.batch.ExcelUpload`` 协议的最小上传对象。"""

    def __init__(self, filename: str, content: bytes):
        self.filename = filename
        self._buf = io.BytesIO(content)

    async def read(self, size: int = -1) -> bytes:
        return self._buf.read(size)


def _make_workbook(sheet_name: str, rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def local_storage(tmp_path, monkeypatch):
    """把对象存储换到 tmp_path，避免往生产的 static/uploads 写测试图片。"""
    monkeypatch.setattr(
        storage_service, '_backend', storage_service.LocalStorageBackend(base_dir=tmp_path),
    )


async def _add_customer(db, name: str) -> Customer:
    """取或建客户：测试库在同一会话内被多个用例共用，同名客户不能重复插。"""
    existing = (await db.execute(select(Customer).where(Customer.name == name))).scalars().first()
    if existing is not None:
        return existing
    customer = Customer(name=name)
    db.add(customer)
    await db.flush()
    return customer


async def _add_product(db, name: str, *, price=0, factory='') -> Product:
    product = Product(name=name, barcode='', price=price, factory=factory,
                      initial_stock=0, current_stock=0)
    db.add(product)
    await db.flush()
    return product


async def _customer_links(db, product_id: int) -> list[int]:
    rows = (await db.execute(
        select(ProductCustomer.customer_id).where(ProductCustomer.product_id == product_id)
    )).scalars().all()
    return list(rows)


# ==================== 模板资产本身 ====================


def test_shipped_template_has_new_sheet_name_and_headers():
    """发布的模板必须已经是新表头 —— 否则用户按模板填的客户/工厂会静默丢弃。"""
    wb = openpyxl.load_workbook(filename=str(TEMPLATE), read_only=True)
    try:
        assert wb.sheetnames == ['新品录入', '订单管理', '备货管理', '出货管理'], wb.sheetnames
        product_ws = wb['新品录入']
        header = [c.value for c in next(product_ws.iter_rows(min_row=1, max_row=1))]
        assert header == HEADERS, f'产品 Sheet 表头不符: {header}'
        stockin_ws = wb['备货管理']
        stockin_header = [c.value for c in next(stockin_ws.iter_rows(min_row=1, max_row=1))]
        assert stockin_header == ['客户名称', '产品名称', '备货数量', '备注'], stockin_header
    finally:
        wb.close()


# ==================== 端点实现（/batch/batch-import） ====================


async def test_import_template_links_customer_and_fills_factory(async_db_session, local_storage):
    """用仓库发布的模板真跑一次：客户归属、工厂、价格三项都要落库。"""
    await _add_customer(async_db_session, 'MINISO国际')
    upload = _Upload(TEMPLATE.name, TEMPLATE.read_bytes())

    result = await batch_import_service.import_products_service(async_db_session, upload)

    assert result.success, result.message
    assert result.imported == 2, result.message
    rows = (await async_db_session.execute(
        select(Product).where(Product.name.like('MINISO长条靠枕%'))
    )).scalars().all()
    assert len(rows) == 2, [p.name for p in rows]
    for product in rows:
        assert float(product.price) > 0, f'{product.name} 价格未落库: {product.price!r}'
        assert product.factory == '东莞', f'{product.name} 工厂未落库: {product.factory!r}'
        assert len(await _customer_links(async_db_session, product.id)) == 1, (
            f'{product.name} 客户归属未落库'
        )


async def test_existing_nonzero_price_is_not_overwritten(async_db_session, local_storage):
    """「只补空值」：已有非 0 价格不能被文件覆盖，空的工厂可以补上。"""
    product = await _add_product(async_db_session, '补值测试-非零价', price=88)
    await _add_customer(async_db_session, 'MINISO国际')
    content = _make_workbook('新品录入', [['MINISO国际', '补值测试-非零价', '湖北', None, 99, None]])

    result = await batch_import_service.import_products_service(
        async_db_session, _Upload('x.xlsx', content))

    await async_db_session.refresh(product)
    assert result.success, result.message
    assert float(product.price) == 88, f'已有价格被覆盖: {product.price!r}'
    assert product.factory == '湖北', f'空工厂未补上: {product.factory!r}'
    assert len(await _customer_links(async_db_session, product.id)) == 1, '空客户归属未补上'


async def test_existing_zero_price_is_filled(async_db_session, local_storage):
    """价格是 0（先建产品、后拿价格）→ 重传文件要能补上。"""
    product = await _add_product(async_db_session, '补值测试-零价', price=0)
    await _add_customer(async_db_session, 'MINISO国际')
    content = _make_workbook('新品录入', [['MINISO国际', '补值测试-零价', '东莞', None, 25, None]])

    result = await batch_import_service.import_products_service(
        async_db_session, _Upload('x.xlsx', content))

    await async_db_session.refresh(product)
    assert result.success, result.message
    assert float(product.price) == 25, f'0 价格未被补上: {product.price!r}'


async def test_existing_product_untouched_reports_not_success(async_db_session, local_storage):
    """字段都已填、又不带图 → 不该再谎报「成功导入」，要说明未改动。"""
    customer = await _add_customer(async_db_session, 'MINISO国际')
    product = await _add_product(async_db_session, '补值测试-已完整', price=9, factory='东莞')
    async_db_session.add(ProductCustomer(product_id=product.id, customer_id=customer.id))
    await async_db_session.flush()
    content = _make_workbook('新品录入', [['MINISO国际', '补值测试-已完整', '东莞', None, 9, None]])

    result = await batch_import_service.import_products_service(
        async_db_session, _Upload('x.xlsx', content))

    assert result.success is False, f'没有任何改动却报成功: {result.message}'
    assert '未改动' in result.message, result.message


async def test_unknown_customer_warns_but_product_still_imported(async_db_session, local_storage):
    """客户名不在主数据：产品照建，但不挂归属，且必须留下可见提示。"""
    content = _make_workbook('新品录入', [['查无此客户', '补值测试-孤儿客户', '东莞', None, 5, None]])

    result = await batch_import_service.import_products_service(
        async_db_session, _Upload('x.xlsx', content))

    product = (await async_db_session.execute(
        select(Product).where(Product.name == '补值测试-孤儿客户'))).scalars().first()
    assert product is not None, '客户不存在不该把产品整行丢掉'
    assert await _customer_links(async_db_session, product.id) == []
    assert '不存在于客户主数据' in result.message, result.message


async def test_legacy_sheet_name_still_recognized(async_db_session, local_storage):
    """旧模板（Sheet 名「产品管理」）继续可导入，不能让存量文件失效。"""
    content = _make_workbook('产品管理', [['MINISO国际', '补值测试-旧Sheet', '东莞', None, 7, None]])

    result = await batch_import_service.import_products_service(
        async_db_session, _Upload('legacy.xlsx', content))

    product = (await async_db_session.execute(
        select(Product).where(Product.name == '补值测试-旧Sheet'))).scalars().first()
    assert result.success, result.message
    assert product is not None and float(product.price) == 7


# ==================== 综合导入 / ARQ 用的 core 实现 ====================


async def test_core_import_also_links_customer_and_factory(async_db_session):
    """core 实现（``import_products_from_data``）必须与端点实现同口径，否则综合导入漏字段。"""
    await _add_customer(async_db_session, 'MINISO国际')
    rows = [{'客户': 'MINISO国际', '产品名称': '补值测试-core', '工厂': '湖北',
             '产品价格': 33, '图片': None}]

    count, messages = await batch_import_service.import_products_from_data(async_db_session, rows)

    product = (await async_db_session.execute(
        select(Product).where(Product.name == '补值测试-core'))).scalars().first()
    assert count == 1, messages
    assert product is not None
    assert float(product.price) == 33 and product.factory == '湖北'
    assert len(await _customer_links(async_db_session, product.id)) == 1


async def test_core_import_backfills_empty_fields_only(async_db_session):
    """core 实现同样「只补空值」：已有价格不动，空工厂补上。"""
    product = await _add_product(async_db_session, '补值测试-core-非零价', price=50)
    await _add_customer(async_db_session, 'MINISO国际')
    rows = [{'客户': 'MINISO国际', '产品名称': '补值测试-core-非零价', '工厂': '东莞',
             '产品价格': 99, '图片': None}]

    await batch_import_service.import_products_from_data(async_db_session, rows)

    await async_db_session.refresh(product)
    assert float(product.price) == 50, f'已有价格被覆盖: {product.price!r}'
    assert product.factory == '东莞'
