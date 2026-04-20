"""
报表统计路由蓝图
包含订单流水、备货流水、出货流水、备货抵扣流水四大功能模块
API路由已迁移至 FastAPI:
  - /api/order-logs → /api/v1/reports/order-logs
  - /api/stockin-logs → /api/v1/reports/stockin-logs
  - /api/stockout-logs → /api/v1/reports/stockout-logs
  - /api/deduction-logs → /api/v1/reports/deduction-logs
此文件仅保留页面渲染路由和导出功能(无FastAPI等效)
"""
from flask import Blueprint, render_template, request, jsonify, session, make_response
from app import db, csrf, limiter
from app.models import StockOut, StockIn, Order, Product
from datetime import datetime, timedelta
from functools import wraps
from sqlalchemy.orm import joinedload
from fastapi_app.shared.report_data_service import get_all_logs_data


def _escape_like(value: str) -> str:
    """转义 SQL LIKE 通配符，防止 LIKE 注入"""
    return value.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
import logging
from utils.report_exporter import ReportExporter, generate_filename
from urllib.parse import quote

logger = logging.getLogger(__name__)

report_bp = Blueprint('report', __name__, url_prefix='/report')


def report_permission_required(permission_code):
    """报表权限验证装饰器"""
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            from flask import flash, redirect, url_for
            if not session.get('user_id'):
                flash('请先登录', 'error')
                return redirect(url_for('main.login'))
            from app.models import User
            user = db.session.get(User, session['user_id'])
            if not user:
                session.clear()
                flash('用户不存在', 'error')
                return redirect(url_for('main.login'))
            if user.is_super:
                return f(*args, **kwargs)
            has_permission = False
            for p in user.permissions:
                if p.code == permission_code:
                    has_permission = True
                    break
            if not has_permission:
                flash('无权限访问该页面', 'error')
                return redirect(url_for('main.index'))
            return f(*args, **kwargs)
        return decorated_function
    return decorator


@report_bp.route('/reports')
@report_permission_required('report_view')
def reports():
    """报表统计页面"""
    active_tab = request.args.get('tab', 'order')
    return render_template('reports.html', active_tab=active_tab)


# ==================== 报表导出功能 (无FastAPI等效) ====================
# get_all_logs_data 已迁移至 fastapi_app.shared.report_data_service，通过顶部 import 使用


@csrf.exempt
@limiter.limit("30 per minute")
@report_bp.route('/api/export', methods=['GET'])
def export_report():
    """
    报表导出API
    参数:
        type: 报表类型 (order, stockin, stockout, deduction)
        customer: 客户筛选 (可选)
        product: 产品筛选 (可选)
        status: 状态筛选 (可选, 仅订单)
        start_date: 开始日期 (可选)
        end_date: 结束日期 (可选)
    """
    try:
        logger.info('开始处理导出请求')

        if not session.get('user_id'):
            logger.warning('用户未登录')
            return jsonify({'success': False, 'error': '请先登录'}), 401

        from app.models import User
        user = db.session.get(User, session['user_id'])
        if not user:
            session.clear()
            logger.warning('用户不存在')
            return jsonify({'success': False, 'error': '用户不存在'}), 401

        has_permission = user.is_super
        if not has_permission:
            for p in user.permissions:
                if p.code == 'report_view':
                    has_permission = True
                    break
        if not has_permission:
            logger.warning('用户无权限')
            return jsonify({'success': False, 'error': '无权限访问'}), 403

        report_type = request.args.get('type', 'order').strip()
        customer = request.args.get('customer', '').strip()
        product = request.args.get('product', '').strip()
        status = request.args.get('status', '').strip()
        start_date = request.args.get('start_date', '').strip()
        end_date = request.args.get('end_date', '').strip()

        logger.info(f'导出参数 - 类型: {report_type}, 客户: {customer}, 产品: {product}')

        valid_types = ['order', 'stockin', 'stockout', 'deduction']
        if report_type not in valid_types:
            logger.warning(f'无效的报表类型: {report_type}')
            return jsonify({'success': False, 'error': '无效的报表类型'}), 400

        logger.info('开始获取数据...')
        logs, stats = get_all_logs_data(
            db.session, report_type, customer, product, status, start_date, end_date
        )

        if not logs:
            logger.info('没有可导出的数据')
            return jsonify({'success': False, 'error': '没有可导出的数据'}), 400

        if len(logs) > 10000:
            logger.warning(f'数据量过大: {len(logs)} 条，建议缩小查询范围')
            return jsonify({'success': False, 'error': f'数据量过大({len(logs)}条)，请缩小查询范围后重试'}), 400

        logger.info(f'获取到 {len(logs)} 条数据，开始生成Excel文件...')
        exporter = ReportExporter()

        if report_type == 'order':
            excel_data = exporter.export_order_logs(logs, stats)
        elif report_type == 'stockin':
            excel_data = exporter.export_stockin_logs(logs, stats)
        elif report_type == 'stockout':
            excel_data = exporter.export_stockout_logs(logs, stats)
        elif report_type == 'deduction':
            excel_data = exporter.export_deduction_logs(logs, stats)
        else:
            return jsonify({'success': False, 'error': '未知的报表类型'}), 400

        filename = generate_filename(report_type)
        logger.info(f'Excel文件生成完成，文件名: {filename}, 大小: {len(excel_data)} 字节')

        encoded_filename = quote(filename)

        response = make_response(excel_data)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f"attachment; filename*=UTF-8''{encoded_filename}"
        response.headers['Content-Length'] = str(len(excel_data))

        logger.info('返回Excel文件响应')
        return response

    except Exception as e:
        logger.error(f'报表导出错误: {str(e)}', exc_info=True)
        return jsonify({'success': False, 'error': '导出失败，请稍后重试'}), 500


@csrf.exempt
@report_bp.route('/api/test-export', methods=['GET'])
def test_export():
    """测试导出接口（需登录）"""
    if 'user_id' not in session:
        return jsonify({'success': False, 'error': '未登录'}), 401
    try:
        logger.info('测试导出接口被调用')

        from openpyxl import Workbook
        from openpyxl.styles import Font
        import io

        wb = Workbook()
        ws = wb.active
        ws.title = '测试报表'

        ws['A1'] = '测试标题'
        ws['A1'].font = Font(bold=True, size=14)

        ws['A2'] = '测试数据1'
        ws['B2'] = '测试数据2'
        ws['A3'] = '值1'
        ws['B3'] = '值2'

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        excel_data = output.getvalue()

        logger.info(f'测试Excel生成成功，大小: {len(excel_data)} 字节')

        response = make_response(excel_data)
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = 'attachment; filename="test_report.xlsx"'
        response.headers['Content-Length'] = str(len(excel_data))

        return response

    except Exception as e:
        logger.error(f'测试导出错误: {str(e)}', exc_info=True)
        return jsonify({'success': False, 'error': '报表导出失败，请稍后重试'}), 500
