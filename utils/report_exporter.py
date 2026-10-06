"""
报表导出工具
支持将报表数据导出为Excel格式
"""
import io
import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logger = logging.getLogger(__name__)


def _sanitize_excel_value(value):
    """清洗Excel公式注入：以=+-@\t\r开头的字符串添加制表符前缀阻止公式解析"""
    if isinstance(value, str) and value and value[0] in '=-+@\t\r':
        return '\t' + value
    return value


class ReportExporter:
    """报表导出器"""

    def __init__(self):
        self.header_font = Font(bold=True, color='FFFFFF')
        self.header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        self.header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.cell_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def export_order_logs(self, logs: list[dict[str, Any]], stats: dict[str, Any]) -> bytes:
        """
        导出订单流水

        Args:
            logs: 订单日志列表
            stats: 统计数据

        Returns:
            Excel文件二进制数据
        """
        wb = Workbook()
        ws = wb.active
        ws.title = '订单流水'

        headers = [
            '下单时间', '客户', '产品名称', '数量', '备注'
        ]

        self._write_sheet(ws, headers, logs, self._format_order_row)

        # 设置订单流水专用的列宽
        ws.column_dimensions['A'].width = 18  # 下单时间
        ws.column_dimensions['B'].width = 12.5  # 客户 2.5cm
        ws.column_dimensions['C'].width = 75  # 产品名称 15cm
        ws.column_dimensions['D'].width = 11  # 数量 2.2cm
        ws.column_dimensions['E'].width = 40  # 备注 8cm

        return self._save_workbook(wb)

    def export_stockin_logs(self, logs: list[dict[str, Any]], stats: dict[str, Any]) -> bytes:
        """
        导出备货流水

        Args:
            logs: 备货日志列表
            stats: 统计数据

        Returns:
            Excel文件二进制数据
        """
        wb = Workbook()
        ws = wb.active
        ws.title = '备货流水'

        headers = [
            '备货时间', '客户', '产品名称', '备货数量', '备注'
        ]

        self._write_sheet(ws, headers, logs, self._format_stockin_row)

        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 12.5
        ws.column_dimensions['C'].width = 75
        ws.column_dimensions['D'].width = 11
        ws.column_dimensions['E'].width = 40

        return self._save_workbook(wb)

    def export_stockout_logs(self, logs: list[dict[str, Any]], stats: dict[str, Any]) -> bytes:
        """
        导出出货流水

        Args:
            logs: 出货日志列表
            stats: 统计数据

        Returns:
            Excel文件二进制数据
        """
        wb = Workbook()
        ws = wb.active
        ws.title = '出货流水'

        headers = [
            '出货时间', '客户', '产品名称', '出货数量', '备注'
        ]

        self._write_sheet(ws, headers, logs, self._format_stockout_row)

        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 12.5
        ws.column_dimensions['C'].width = 75
        ws.column_dimensions['D'].width = 11
        ws.column_dimensions['E'].width = 40

        return self._save_workbook(wb)

    def export_deduction_logs(self, logs: list[dict[str, Any]], stats: dict[str, Any]) -> bytes:
        """
        导出备货抵扣流水

        Args:
            logs: 抵扣日志列表
            stats: 统计数据

        Returns:
            Excel文件二进制数据
        """
        wb = Workbook()
        ws = wb.active
        ws.title = '备货抵扣流水'

        headers = [
            '抵扣时间', '客户', '产品名称', '抵扣数量', '订单备注'
        ]

        self._write_sheet(ws, headers, logs, self._format_deduction_row)

        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 12.5
        ws.column_dimensions['C'].width = 75
        ws.column_dimensions['D'].width = 11
        ws.column_dimensions['E'].width = 40

        return self._save_workbook(wb)

    def export_return_logs(self, logs: list[dict[str, Any]], stats: dict[str, Any]) -> bytes:
        """
        导出退货流水

        Args:
            logs: 退货日志列表
            stats: 统计数据

        Returns:
            Excel文件二进制数据
        """
        wb = Workbook()
        ws = wb.active
        ws.title = '退货流水'

        headers = [
            '退货时间', '客户', '产品名称', '退货数量', '备注'
        ]

        self._write_sheet(ws, headers, logs, self._format_return_row)

        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 12.5
        ws.column_dimensions['C'].width = 75
        ws.column_dimensions['D'].width = 11
        ws.column_dimensions['E'].width = 40

        return self._save_workbook(wb)

    def export_inventory_data(self, inventory_data: list[dict[str, Any]], stats: dict[str, Any]) -> bytes:
        """
        导出库存汇总数据

        Args:
            inventory_data: 库存数据列表
            stats: 统计数据

        Returns:
            Excel文件二进制数据
        """
        wb = Workbook()
        ws = wb.active
        ws.title = '库存汇总'

        headers = ['客户', '产品名称', '订单数量', '备货数量', '总库存']

        self._write_inventory_sheet(ws, headers, inventory_data)

        return self._save_workbook(wb)

    def _write_inventory_sheet(self, ws, headers, inventory_data):
        """写入库存汇总工作表"""
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.thin_border

        for row_idx, item in enumerate(inventory_data, 2):
            row_data = [
                item.get('customer', ''), item.get('product_name', ''),
                item.get('order_quantity', 0),
                item.get('stockin_quantity', 0), item.get('total_stock', 0),
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=_sanitize_excel_value(value))
                cell.alignment = self.cell_alignment
                cell.border = self.thin_border

        # 设置默认行高
        for row_idx in range(1, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 28.35

        # 设置列宽
        for col_idx in range(1, len(headers) + 1):
            col_letter = chr(64 + col_idx) if col_idx <= 26 else chr(64 + (col_idx - 1) // 26) + chr(65 + (col_idx - 1) % 26)
            if col_idx == 2:  # B列 产品名称
                ws.column_dimensions[col_letter].width = 75
            elif headers[col_idx - 1] in ('备货数量', '总库存', '订单数量'):
                ws.column_dimensions[col_letter].width = 12.5
            else:
                ws.column_dimensions[col_letter].width = 18

    def _format_inventory_row(self, log: dict) -> list:
        """格式化库存行数据"""
        return [
            log.get('customer', ''),
            log.get('product_name', ''),
            log.get('order_quantity', 0),
            log.get('stockin_quantity', 0),
            log.get('total_stock', 0)
        ]

    def _write_sheet(self, ws, headers: list[str], logs: list[dict], row_formatter):
        """
        写入工作表内容

        Args:
            ws: 工作表
            headers: 表头列表
            logs: 数据列表
            row_formatter: 行数据格式化函数
        """
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.thin_border

        for row_idx, log in enumerate(logs, 2):
            row_data = row_formatter(log)
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=_sanitize_excel_value(value))
                cell.alignment = self.cell_alignment
                cell.border = self.thin_border

        # 设置默认行高为1cm (约28.35磅)
        for row_idx in range(1, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 28.35

        # 设置列宽
        # 1cm约等于5个字符宽度
        for col_idx in range(1, len(headers) + 1):
            col_letter = chr(64 + col_idx)
            if col_idx == 2:  # B列
                ws.column_dimensions[col_letter].width = 75  # 15cm
            elif col_idx in [3, 4, 5]:  # C、D、E列
                ws.column_dimensions[col_letter].width = 12.5  # 2.5cm
            else:
                ws.column_dimensions[col_letter].width = 18

    def _format_order_row(self, log: dict) -> list:
        """格式化订单行数据"""
        return [
            log.get('order_date', ''),
            log.get('customer', ''),
            log.get('product_name', ''),
            log.get('quantity', 0),
            log.get('note', '')
        ]

    def _format_stockin_row(self, log: dict) -> list:
        """格式化备货行数据"""
        return [
            log.get('stockin_date', ''),
            log.get('customer', ''),
            log.get('product_name', ''),
            log.get('quantity', 0),
            log.get('note', '')
        ]

    def _format_stockout_row(self, log: dict) -> list:
        """格式化出货行数据"""
        return [
            log.get('stockout_date', ''),
            log.get('customer', ''),
            log.get('product_name', ''),
            log.get('quantity', 0),
            log.get('note', '')
        ]

    def _format_deduction_row(self, log: dict) -> list:
        """格式化抵扣行数据"""
        return [
            log.get('deduction_time', ''),
            log.get('customer', ''),
            log.get('product_name', ''),
            log.get('quantity', 0),
            log.get('order_note', '')
        ]

    def _format_return_row(self, log: dict) -> list:
        """格式化退货行数据"""
        return [
            log.get('return_date', ''),
            log.get('customer', ''),
            log.get('product_name', ''),
            log.get('quantity', 0),
            log.get('note', '')
        ]

    def _write_summary(self, ws, stats: dict, title: str, col_count: int):
        """
        写入统计信息

        Args:
            ws: 工作表
            stats: 统计数据
            title: 统计标题
            col_count: 列数
        """
        start_row = ws.max_row + 3

        title_cell = ws.cell(row=start_row, column=1, value=title)
        title_cell.font = Font(bold=True, size=14)
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=col_count)

        row_idx = start_row + 1
        for key, value in stats.items():
            key_cell = ws.cell(row=row_idx, column=1, value=self._format_stat_key(key))
            key_cell.font = Font(bold=True)
            key_cell.border = self.thin_border

            value_cell = ws.cell(row=row_idx, column=2, value=value)
            value_cell.border = self.thin_border
            row_idx += 1

    def _format_stat_key(self, key: str) -> str:
        """格式化统计键名"""
        key_map = {
            'total_count': '总记录数',
            'total_quantity': '总数量',
            'pending_count': '待处理数',
            'completed_count': '已完成数',
            'total_remaining': '总剩余数量',
            'total_orders': '涉及订单数'
        }
        return key_map.get(key, key)

    def _save_workbook(self, wb) -> bytes:
        """
        保存工作簿到内存

        Args:
            wb: 工作簿对象

        Returns:
            二进制数据
        """
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()


def generate_filename(report_type: str) -> str:
    """
    生成导出文件名

    Args:
        report_type: 报表类型

    Returns:
        文件名
    """
    type_map = {
        'order': '订单流水',
        'stockin': '备货流水',
        'stockout': '出货流水',
        'deduction': '备货抵扣流水',
        'return': '退货流水',
        'inventory': '库存汇总'
    }
    name = type_map.get(report_type, '报表')
    return f'{name}.xlsx'
