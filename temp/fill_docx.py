# -*- coding: utf-8 -*-
"""将采购计划 xlsx 的 A 列(sku名称)、B 列(条码) 填入 符合性声明 docx 的 Annex 表格 BarCode/第二代码 列。
映射：B列(条码) -> BarCode 行，A列(sku名称) -> 第二代码 行。
表格从模板 4 行扩展到全部数据行，保留原表格结构与样式。
"""
import glob
import os
from copy import deepcopy

import docx
import openpyxl
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

XLSX = glob.glob(r"C:\Users\*\AppData\Local\Temp\codebuddy-dropped-files\*\采购计划视图*.xlsx")[0]
DOCX = glob.glob(r"C:\Users\*\AppData\Local\Temp\codebuddy-dropped-files\*\符合性声明*.docx")[0]
OUT = os.path.join(os.path.dirname(DOCX), "符合性声明DOC-PPWR-V1.0-已填写.docx")

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# ---------- 1. 读取 Excel A/B 列 ----------
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb.active
rows_data = []
for r in ws.iter_rows(min_row=2):
    a = r[0].value
    b = r[1].value
    if a is None and b is None:
        continue
    rows_data.append((str(a).strip() if a is not None else "", str(b).strip() if b is not None else ""))
print(f"Excel 数据行数: {len(rows_data)}")

# ---------- 2. 打开 docx，扩展 Annex 表格数据行 ----------
d = docx.Document(DOCX)
tbl = d.tables[0]._tbl
trs = tbl.findall(W + "tr")

HEADER_ROW = 13
FIRST_DATA_ROW = 14
LAST_DATA_ROW = 17   # 模板已有的 4 行 (编号 1-4)
AFTER_DATA_ROW = 18  # Date of issue 行

template_tr = trs[FIRST_DATA_ROW]
after_tr = trs[AFTER_DATA_ROW]

need = len(rows_data)
existing = LAST_DATA_ROW - FIRST_DATA_ROW + 1  # 4
if need < existing:
    raise SystemExit(f"数据行({need})少于模板行({existing})，无需扩展")

# 以模板行 deepcopy 插入，保留全部格式（边框、合并、行高、居中）
for _ in range(need - existing):
    after_tr.addprevious(deepcopy(template_tr))

# ---------- 3. 填充每行数据 ----------
trs2 = tbl.findall(W + "tr")
data_trs = trs2[FIRST_DATA_ROW : FIRST_DATA_ROW + need]


def make_run(text, bold=False):
    """构造与模板格式一致的 run：Arial/宋体 10pt"""
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rFonts.set(qn("w:eastAsia"), "宋体")
    rPr.append(rFonts)
    if bold:
        rPr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "20")
    rPr.append(szCs)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def make_para(runs, center=True):
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    if center:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        pPr.append(jc)
    p.append(pPr)
    for run in runs:
        p.append(run)
    return p


def set_tc(tc, paras):
    for p in tc.findall(W + "p"):
        tc.remove(p)
    for para in paras:
        tc.append(para)


for idx, (sku, barcode) in enumerate(rows_data, start=1):
    tcs = data_trs[idx - 1].findall(W + "tc")
    # tc0: No. 编号
    set_tc(tcs[0], [make_para([make_run(str(idx))])])
    # tc1: BarCode/第二代码 单元格 -> 第一行条码(BarCode)，第二行商品名(第二代码)
    set_tc(
        tcs[1],
        [
            make_para([make_run(barcode)]),
            make_para([make_run(sku)]),
        ],
    )

d.save(OUT)
print("已保存:", OUT)

# ---------- 4. 验证 ----------
d2 = docx.Document(OUT)
tb2 = d2.tables[0]
trs3 = tb2._tbl.findall(W + "tr")
print(f"输出表格总行数: {len(trs3)} (原 21 + 新增 {len(trs3) - 21})")
for i in (0, 1, len(rows_data) - 1):
    tcs = trs3[FIRST_DATA_ROW + i].findall(W + "tc")
    no = tcs[0].find(W + "p").find(W + "r").find(W + "t").text
    paras = tcs[1].findall(W + "p")
    lines = ["".join(t.text or "" for t in p.findall(".//" + W + "t")) for p in paras]
    print(f"  数据行 #{i+1}: No={no!r} | BarCode/第二代码={lines}")
